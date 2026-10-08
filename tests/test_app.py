import io
from datetime import datetime, timedelta, timezone
import json
import pytest
from app import create_app

@pytest.fixture()
def app(tmp_path):
    return create_app({"TESTING": True, "DATABASE": str(tmp_path / "test.sqlite"), "SEED_DEMO": False, "SECRET_KEY": "test"})

@pytest.fixture()
def client(app):
    return app.test_client()

def create_panel(client, panel_id="CMP1000"):
    return client.post("/panels/new", data={
        "panel_id": panel_id,
        "panel_name": "Test Transformers",
        "category": "Transformers",
        "business": "GI",
        "region_level": "Country",
        "region_value": "United Kingdom",
        "owner": "Test Owner",
        "mdf_codes": ["MDF-TR-001"],
        "lead_mdf_code": "MDF-TR-001",
        "fields_json": json.dumps([{"fieldId":"rating","fieldName":"Rating","type":"number","required":False,"options":[]}]),
    })

def test_empty_portfolio(client):
    response=client.get("/")
    assert response.status_code==200
    assert b"Category panel" in response.data

def test_create_panel_and_api(client):
    response=create_panel(client)
    assert response.status_code==302
    api=client.get("/api/panels/CMP1000")
    assert api.status_code==200
    body=api.get_json()
    assert body["panel"]["panelName"]=="Test Transformers"
    assert body["panel"]["region"]["value"]=="United Kingdom"
    assert body["panel"]["supplierFields"][0]["fieldId"]=="rating"

def test_edit_panel_excludes_custom_fields_and_configuration_owns_them(client):
    create_panel(client)
    edit_page=client.get("/panels/CMP1000/edit")
    assert edit_page.status_code==200
    assert b"Panel details" in edit_page.data
    assert b"Supplier data definition" not in edit_page.data
    assert b"Rating" not in edit_page.data

    config_page=client.get("/panels/CMP1000/configuration")
    assert config_page.status_code==200
    assert b"Panel Configuration" in config_page.data
    assert b"Supplier data definition" in config_page.data
    assert b"Rating" in config_page.data

    view=client.get("/panels/CMP1000")
    assert b"Panel Configuration" in view.data


def test_duplicate_panel_id_rejected(client):
    create_panel(client)
    response=create_panel(client)
    assert response.status_code==409


def test_custom_field_groups_render_on_supplier_and_bulk_forms(client):
    response=client.post("/panels/new", data={
        "panel_id":"CMP-GROUP",
        "panel_name":"Grouped Panel",
        "category":"Transformers",
        "business":"GI",
        "region_level":"Country",
        "region_value":"United Kingdom",
        "owner":"Test Owner",
        "mdf_codes":["MDF-TR-001"],
        "lead_mdf_code":"MDF-TR-001",
        "field_groups_json":json.dumps([
            {"groupId":"risk_qualification","name":"Risk & Qualification","order":1},
            {"groupId":"performance","name":"Performance","order":2},
        ]),
        "fields_json":json.dumps([
            {"fieldId":"supplier_risk","fieldName":"Supplier Risk","type":"text","required":False,"options":[],"groupId":"risk_qualification"},
            {"fieldId":"quality","fieldName":"Quality","type":"number","required":False,"options":[],"groupId":"performance"},
            {"fieldId":"notes","fieldName":"Notes","type":"text","required":False,"options":[]},
        ]),
    })
    assert response.status_code==302
    body=client.get("/api/panels/CMP-GROUP").get_json()
    assert body["schemaVersion"]==5
    assert [g["name"] for g in body["panel"]["fieldGroups"]]==["Risk & Qualification","Performance"]
    assert body["panel"]["supplierFields"][0]["groupId"]=="risk_qualification"
    assert body["panel"]["supplierFields"][2]["groupId"]==""

    supplier_form=client.get("/panels/CMP-GROUP/suppliers/new")
    assert supplier_form.status_code==200
    assert b"Risk &amp; Qualification" in supplier_form.data
    assert b"Performance" in supplier_form.data
    assert b"Ungrouped" in supplier_form.data

    client.post("/panels/CMP-GROUP/suppliers/new", data={
        "supplier_id":"SUP-G1","supplier_name":"Grouped Supplier","address":"","post_code":"",
        "custom_supplier_risk":"Low","custom_quality":"4","custom_notes":"Legacy-compatible",
    })
    bulk=client.get("/panels/CMP-GROUP/suppliers/bulk-edit")
    assert bulk.status_code==200
    assert b"Risk &amp; Qualification" in bulk.data
    assert b"Performance" in bulk.data
    assert b"Ungrouped" in bulk.data


def test_existing_schema3_panel_migrates_without_losing_supplier_values(app):
    client=app.test_client()
    create_panel(client, "CMP-OLD")
    client.post("/panels/CMP-OLD/suppliers/new", data={
        "supplier_id":"SUP-OLD","supplier_name":"Existing Supplier","address":"","post_code":"",
        "custom_rating":"275",
    })
    with app.app_context():
        import sqlite3
        db=sqlite3.connect(app.config["DATABASE"])
        row=db.execute("SELECT data_json FROM panels WHERE panel_id='CMP-OLD'").fetchone()
        payload=json.loads(row[0])
        payload["schemaVersion"]=3
        payload["panel"].pop("fieldGroups", None)
        for field in payload["panel"]["supplierFields"]:
            field.pop("groupId", None)
        payload["panel"]["metadata"]["version"]=3
        db.execute("UPDATE panels SET data_json=? WHERE panel_id='CMP-OLD'", (json.dumps(payload),))
        db.commit()
        db.close()

    migrated_app=create_app({
        "TESTING":True,
        "DATABASE":app.config["DATABASE"],
        "SEED_DEMO":False,
        "SECRET_KEY":"test",
    })
    migrated=migrated_app.test_client().get("/api/panels/CMP-OLD").get_json()
    assert migrated["schemaVersion"]==5
    assert migrated["panel"]["fieldGroups"]==[]
    assert migrated["panel"]["supplierFields"][0]["groupId"]==""
    assert migrated["panel"]["suppliers"][0]["customFields"]["rating"]==275.0
    assert "schema4MigratedAt" in migrated["panel"]["metadata"]

def test_add_supplier_uses_custom_schema(client):
    create_panel(client)
    response=client.post("/panels/CMP1000/suppliers/new", data={"supplier_id":"SUP-1","supplier_name":"Supplier One","address":"Road","post_code":"AA1 1AA","custom_rating":"250"})
    assert response.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    supplier=body["panel"]["suppliers"][0]
    assert supplier["supplierId"]=="SUP-1"
    assert supplier["customFields"]["rating"]==250.0

def test_raw_json_rejects_panel_id_change(client):
    create_panel(client)
    payload=client.get("/api/panels/CMP1000").get_json()
    payload["panel"]["panelId"]="OTHER"
    response=client.post("/panels/CMP1000/data", data={"raw_json":json.dumps(payload)})
    assert response.status_code==400


def test_edit_supplier_preserves_id_and_records_audit(client):
    create_panel(client)
    client.post("/panels/CMP1000/suppliers/new", data={
        "supplier_id":"SUP-1","supplier_name":"Supplier One","address":"Road","post_code":"AA1 1AA","custom_rating":"250"
    })
    response=client.post("/panels/CMP1000/suppliers/SUP-1/edit", data={
        "supplier_id":"SUP-1","supplier_name":"Supplier Updated","address":"New Road","post_code":"BB2 2BB","custom_rating":"300"
    })
    assert response.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    supplier=body["panel"]["suppliers"][0]
    assert supplier["supplierId"]=="SUP-1"
    assert supplier["supplierName"]=="Supplier Updated"
    assert supplier["customFields"]["rating"]==300.0
    audit=body["panel"]["metadata"]["auditTrail"]
    assert audit[-1]["action"]=="supplier_updated"
    assert audit[-1]["details"]["before"]["supplierName"]=="Supplier One"

def test_delete_supplier_requires_exact_confirmation_and_keeps_snapshot(client):
    create_panel(client)
    client.post("/panels/CMP1000/suppliers/new", data={
        "supplier_id":"SUP-1","supplier_name":"Supplier One","address":"Road","post_code":"AA1 1AA","custom_rating":"250"
    })
    refused=client.post("/panels/CMP1000/suppliers/SUP-1/delete", data={"confirm_supplier_id":"wrong"})
    assert refused.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    assert len(body["panel"]["suppliers"])==1

    accepted=client.post("/panels/CMP1000/suppliers/SUP-1/delete", data={"confirm_supplier_id":"SUP-1"})
    assert accepted.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    assert body["panel"]["suppliers"]==[]
    audit=body["panel"]["metadata"]["auditTrail"]
    assert audit[-1]["action"]=="supplier_deleted"
    assert audit[-1]["details"]["snapshot"]["supplierId"]=="SUP-1"


def test_archive_panel_requires_confirmation_and_is_reversible(client):
    create_panel(client)
    refused=client.post("/panels/CMP1000/archive", data={"confirm_panel_id":"wrong"})
    assert refused.status_code==302
    assert client.get("/api/panels/CMP1000").status_code==200
    assert b"CMP1000" in client.get("/").data

    archived=client.post("/panels/CMP1000/archive", data={"confirm_panel_id":"CMP1000"})
    assert archived.status_code==302
    active_page=client.get("/").data
    archived_page=client.get("/?archived=1").data
    assert b'href="/panels/CMP1000"' not in active_page
    assert b'href="/panels/CMP1000"' in archived_page
    body=client.get("/api/panels/CMP1000").get_json()
    assert body["panel"]["metadata"]["status"]=="archived"
    assert body["panel"]["metadata"]["auditTrail"][-1]["action"]=="panel_archived"

    restored=client.post("/panels/CMP1000/restore")
    assert restored.status_code==302
    assert b"CMP1000" in client.get("/").data
    body=client.get("/api/panels/CMP1000").get_json()
    assert body["panel"]["metadata"]["status"]=="active"
    assert body["panel"]["metadata"]["auditTrail"][-1]["action"]=="panel_restored"


def test_mdf_master_data_admin_and_usage_guard(client):
    page=client.get("/settings/mdf")
    assert page.status_code==200
    assert b"MDF-TR-001" in page.data

    added=client.post("/settings/mdf", data={"code":"MDF-NEW-001","description":"New Equipment"})
    assert added.status_code==302
    page=client.get("/settings/mdf")
    assert b"MDF-NEW-001" in page.data
    assert b"New Equipment" in page.data

    create_panel(client)
    blocked=client.post("/settings/mdf/MDF-TR-001/toggle")
    assert blocked.status_code==302
    page=client.get("/settings/mdf")
    assert b"used by 1 active panel" in page.data

    deactivated=client.post("/settings/mdf/MDF-NEW-001/toggle")
    assert deactivated.status_code==302
    new_panel=client.get("/panels/new")
    assert b'value="MDF-NEW-001"' not in new_panel.data

    reactivated=client.post("/settings/mdf/MDF-NEW-001/toggle")
    assert reactivated.status_code==302
    new_panel=client.get("/panels/new")
    assert b'value="MDF-NEW-001"' in new_panel.data


def test_panel_supports_multiple_mdf_codes(client):
    response=client.post("/panels/new", data={
        "panel_id":"CMP2000",
        "panel_name":"Multi MDF Panel",
        "category":"Transformers",
        "business":"GI",
        "region_level":"Global",
        "region_value":"Global",
        "owner":"Owner",
        "mdf_codes":["MDF-TR-001","3GF"],
        "lead_mdf_code":"MDF-TR-001",
        "fields_json":"[]",
    })
    assert response.status_code==302
    body=client.get("/api/panels/CMP2000").get_json()
    assert body["panel"]["leadMdfCode"]=="MDF-TR-001"
    assert set(body["panel"]["mdfCodes"])=={"MDF-TR-001","3GF"}
    assert body["panel"]["mdfCode"]=="MDF-TR-001"

def test_lead_mdf_must_be_selected(client):
    response=client.post("/panels/new", data={
        "panel_id":"CMP2001",
        "panel_name":"Invalid MDF Panel",
        "category":"Transformers",
        "business":"GI",
        "region_level":"Global",
        "region_value":"Global",
        "owner":"Owner",
        "mdf_codes":["3GF"],
        "lead_mdf_code":"MDF-TR-001",
        "fields_json":"[]",
    })
    assert response.status_code==400


def test_new_custom_field_can_be_amended_by_supplier_form_and_bulk_editor(client):
    create_panel(client)
    client.post("/panels/CMP1000/suppliers/new", data={
        "supplier_id":"SUP-1",
        "supplier_name":"Supplier One",
        "address":"Road",
        "post_code":"AA1 1AA",
        "custom_rating":"250",
    })

    updated_panel=client.post("/panels/CMP1000/configuration", data={
        "field_groups_json":"[]",
        "fields_json":json.dumps([
            {"fieldId":"rating","fieldName":"Rating","type":"number","required":False,"options":[]},
            {"fieldId":"review_status","fieldName":"Review Status","type":"dropdown","required":False,"options":["Open","Closed"]},
        ]),
    })
    assert updated_panel.status_code==302

    supplier_form=client.get("/panels/CMP1000/suppliers/SUP-1/edit")
    assert supplier_form.status_code==200
    assert b"Review Status" in supplier_form.data

    panel_view=client.get("/panels/CMP1000")
    assert b'/panels/CMP1000/suppliers/SUP-1/edit' in panel_view.data
    assert b"Edit supplier data" in panel_view.data

    bulk_page=client.get("/panels/CMP1000/suppliers/bulk-edit")
    assert bulk_page.status_code==200
    assert b"Review Status" in bulk_page.data
    assert b"Supplier One" in bulk_page.data

    saved=client.post("/panels/CMP1000/suppliers/bulk-edit", data={
        "SUP-1__rating":"275",
        "SUP-1__review_status":"Open",
    })
    assert saved.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    supplier=body["panel"]["suppliers"][0]
    assert supplier["supplierName"]=="Supplier One"
    assert supplier["address"]=="Road"
    assert supplier["postCode"]=="AA1 1AA"
    assert supplier["customFields"]["rating"]==275.0
    assert supplier["customFields"]["review_status"]=="Open"
    assert body["panel"]["metadata"]["auditTrail"][-1]["action"]=="supplier_bulk_custom_fields_updated"


def test_mdf_selector_is_searchable_compact_multiselect(client):
    page=client.get("/panels/new")
    assert page.status_code==200
    assert b"mdf-select-trigger" in page.data
    assert b"Search MDF code or description" in page.data
    assert b"mdf-check-grid" not in page.data
    assert b'value="3AA"' in page.data
    assert b'value="5BZ"' in page.data

def test_uploaded_mdf_catalogue_is_loaded(client):
    page=client.get("/settings/mdf")
    assert page.status_code==200
    assert b"Terminal Blocks" in page.data
    assert b"Offshore Route Preparation Services" in page.data


def test_portfolio_favourites_sorting_pagination_and_persistent_filters(client):
    for idx in range(12):
        response=client.post("/panels/new", data={
            "panel_id":f"CMP3{idx:03d}",
            "panel_name":f"Panel {chr(65 + (11-idx))}",
            "category":"Transformers" if idx % 2 == 0 else "Switchgear",
            "business":"GI" if idx % 2 == 0 else "GPQSS",
            "region_level":"Country",
            "region_value":"United Kingdom",
            "owner":f"Owner {idx:02d}",
            "mdf_codes":["MDF-TR-001"],
            "lead_mdf_code":"MDF-TR-001",
            "fields_json":"[]",
        })
        assert response.status_code==302

    fav=client.post("/panels/CMP3000/favourite", data={"return_to":"/"})
    assert fav.status_code==302
    favourites=client.get("/?favourites=1")
    assert b"CMP3000" in favourites.data
    assert b"CMP3001" not in favourites.data

    page=client.get("/?category=Transformers&business=GI&q=Panel&page_size=10&sort=panel&direction=asc")
    assert page.status_code==200
    assert b"Page 1 of" in page.data
    assert b"10 rows" in page.data
    first_pos=page.data.find(b"Panel B")
    later_pos=page.data.find(b"Panel D")
    assert first_pos != -1 and later_pos != -1 and first_pos < later_pos

    with client.session_transaction() as sess:
        saved=sess["portfolio_filters"]
        assert saved["category"]=="Transformers"
        assert saved["business"]=="GI"
        assert saved["q"]=="Panel"
        assert saved["sort"]=="panel"
        assert saved["direction"]=="asc"

    persisted=client.get("/")
    assert b'value="Panel"' in persisted.data
    assert b"Transformers" in persisted.data

    reset=client.get("/?clear=1", follow_redirects=True)
    assert reset.status_code==200
    with client.session_transaction() as sess:
        assert "portfolio_filters" not in sess

def test_portfolio_paginates_large_result_sets(client):
    for idx in range(11):
        response=client.post("/panels/new", data={
            "panel_id":f"CMP4{idx:03d}",
            "panel_name":f"Paged Panel {idx:02d}",
            "category":"Transformers",
            "business":"GI",
            "region_level":"Global",
            "region_value":"Global",
            "owner":"Owner",
            "mdf_codes":["MDF-TR-001"],
            "lead_mdf_code":"MDF-TR-001",
            "fields_json":"[]",
        })
        assert response.status_code==302
    first=client.get("/?q=Paged+Panel&page_size=10&sort=panel&direction=asc")
    second=client.get("/?q=Paged+Panel&page_size=10&sort=panel&direction=asc&page=2")
    assert b"Page 1 of 2" in first.data
    assert b"Page 2 of 2" in second.data
    assert b"CMP4010" in second.data


def test_supplier_master_import_generates_addresses_and_supports_typeahead(client):
    csv_data = b"BPID,Supplier_Name\n1000000001,Alpha Engineering\n1000000002,Beta Power Ltd\n"
    response = client.post(
        "/settings/suppliers",
        data={"supplier_file": (io.BytesIO(csv_data), "suppliers.csv")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 302

    search = client.get("/api/supplier-master/search?q=Alpha")
    assert search.status_code == 200
    rows = search.get_json()
    assert rows[0]["bpid"] == "1000000001"
    assert rows[0]["supplier_name"] == "Alpha Engineering"
    assert rows[0]["address"]
    assert rows[0]["post_code"]
    assert rows[0]["address_source"] == "generated"

    by_bpid = client.get("/api/supplier-master/search?q=1000000002")
    assert by_bpid.get_json()[0]["supplier_name"] == "Beta Power Ltd"

    form = client.get("/panels/new")
    assert form.status_code == 200

def test_supplier_master_remove_hides_from_selection_but_can_restore(client):
    csv_data = b"BPID,Supplier_Name\n1000000010,Remove Me Ltd\n"
    client.post(
        "/settings/suppliers",
        data={"supplier_file": (io.BytesIO(csv_data), "suppliers.csv")},
        content_type="multipart/form-data",
    )
    assert client.get("/api/supplier-master/search?q=Remove").get_json()

    removed = client.post("/settings/suppliers/1000000010/remove")
    assert removed.status_code == 302
    assert client.get("/api/supplier-master/search?q=Remove").get_json() == []

    removed_page = client.get("/settings/suppliers?removed=1&q=Remove")
    assert b"Remove Me Ltd" in removed_page.data
    assert b"Removed" in removed_page.data

    restored = client.post("/settings/suppliers/1000000010/restore")
    assert restored.status_code == 302
    assert client.get("/api/supplier-master/search?q=Remove").get_json()

def test_add_supplier_uses_master_data_when_bpid_selected(client):
    create_panel(client)
    csv_data = b"BPID,Supplier_Name,Address,Post_Code\n1000000020,Master Supplier Ltd,1 Real Road,AB1 2CD\n"
    client.post(
        "/settings/suppliers",
        data={"supplier_file": (io.BytesIO(csv_data), "suppliers.csv")},
        content_type="multipart/form-data",
    )
    response = client.post("/panels/CMP1000/suppliers/new", data={
        "supplier_id":"1000000020",
        "supplier_name":"Wrong Manual Name",
        "address":"Wrong Address",
        "post_code":"ZZ1 1ZZ",
        "custom_rating":"123",
    })
    assert response.status_code == 302
    supplier = client.get("/api/panels/CMP1000").get_json()["panel"]["suppliers"][0]
    assert supplier["supplierName"] == "Master Supplier Ltd"
    assert supplier["address"] == "1 Real Road"
    assert supplier["postCode"] == "AB1 2CD"
    assert supplier["masterLinked"] is True

def test_supplier_form_contains_master_typeahead(client):
    create_panel(client)
    page = client.get("/panels/CMP1000/suppliers/new")
    assert page.status_code == 200
    assert b"Find supplier by BPID or name" in page.data
    assert b"supplier-master-search" in page.data


def test_dashboard_uses_configured_currency_and_review_semantics(client):
    fields = [
        {"fieldId":"annual_spend","fieldName":"Annual Spend","type":"number","required":False,"options":[]},
        {"fieldId":"spend_currency","fieldName":"Spend Currency","type":"dropdown","required":False,"options":["GBP","EUR","USD"]},
        {"fieldId":"qualification_status","fieldName":"Qualification Status","type":"dropdown","required":False,"options":["Qualified","In review"]},
        {"fieldId":"qualification_review_date","fieldName":"Qualification Review Date","type":"date","required":False,"options":[]},
        {"fieldId":"classification","fieldName":"Classification","type":"dropdown","required":False,"options":["Preferred","Standard"]},
    ]
    response=client.post("/panels/new", data={
        "panel_id":"CMP5000",
        "panel_name":"Dashboard Panel",
        "category":"Transformers",
        "business":"GI",
        "region_level":"Global",
        "region_value":"Global",
        "owner":"Owner",
        "mdf_codes":["MDF-TR-001"],
        "lead_mdf_code":"MDF-TR-001",
        "fields_json":json.dumps(fields),
    })
    assert response.status_code==302

    today=datetime.now(timezone.utc).date()
    suppliers=[
        ("SUP-D1","Alpha Supplier","1000000","GBP","Qualified",(today-timedelta(days=1)).isoformat(),"Preferred"),
        ("SUP-D2","Alpha Supplier","2000000","EUR","In review",(today+timedelta(days=10)).isoformat(),"Standard"),
        ("SUP-D3","Gamma Supplier","3000000","USD","Qualified","", "Preferred"),
    ]
    for sid,name,spend,currency,qual,review,classification in suppliers:
        response=client.post("/panels/CMP5000/suppliers/new", data={
            "supplier_id":sid,
            "supplier_name":name,
            "address":"Road",
            "post_code":"AA1 1AA",
            "custom_annual_spend":spend,
            "custom_spend_currency":currency,
            "custom_qualification_status":qual,
            "custom_qualification_review_date":review,
            "custom_classification":classification,
        })
        assert response.status_code==302

    settings=client.post("/settings/dashboard", data={
        "base_currency":"GBP",
        "spend_field_id":"annual_spend",
        "currency_field_id":"spend_currency",
        "qualification_field_id":"qualification_status",
        "qualification_review_field_id":"qualification_review_date",
        "classification_field_id":"classification",
        "currency_rates":"GBP=1\nEUR=0.5",
    })
    assert settings.status_code==302

    page=client.get("/")
    assert page.status_code==200
    assert b"GBP 2.0M" in page.data
    assert b"USD 3,000,000" in page.data
    assert b"Alpha Supplier" in page.data
    assert b"2</strong>" in page.data
    assert b"1 overdue" in page.data
    assert b"Due in 30 days" in page.data
    assert b"No review date" in page.data

def test_dashboard_settings_reject_invalid_currency_rates(client):
    response=client.post("/settings/dashboard", data={
        "base_currency":"GBP",
        "spend_field_id":"annual_spend",
        "currency_field_id":"spend_currency",
        "qualification_field_id":"qualification_status",
        "qualification_review_field_id":"qualification_review_date",
        "classification_field_id":"classification",
        "currency_rates":"GBP=1\nEUR=not-a-number",
    })
    assert response.status_code==400
    assert b"Currency rates must use" in response.data


def test_field_removal_previews_impact_preserves_orphan_and_allows_restore_or_purge(client):
    create_panel(client)
    client.post("/panels/CMP1000/suppliers/new", data={
        "supplier_id":"SUP-ORPHAN",
        "supplier_name":"Orphan Supplier",
        "address":"Road",
        "post_code":"AA1 1AA",
        "custom_rating":"250",
    })
    edit_data={"field_groups_json":"[]","fields_json":"[]"}

    preview=client.post("/panels/CMP1000/configuration", data=edit_data)
    assert preview.status_code==409
    assert b"Field removal impact" in preview.data
    assert b"Rating" in preview.data
    assert b"SUP-ORPHAN" in preview.data

    before=client.get("/api/panels/CMP1000").get_json()
    assert before["panel"]["supplierFields"][0]["fieldId"]=="rating"
    assert before["panel"]["suppliers"][0]["customFields"]["rating"]==250.0

    confirmed=client.post("/panels/CMP1000/configuration", data={**edit_data,"confirm_field_removal":"1"})
    assert confirmed.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    assert body["panel"]["supplierFields"]==[]
    assert body["panel"]["suppliers"][0]["customFields"]["rating"]==250.0
    orphan=body["panel"]["metadata"]["orphanedSupplierFields"]["rating"]
    assert orphan["field"]["fieldName"]=="Rating"
    assert orphan["affectedSuppliers"]==1

    orphan_page=client.get("/panels/CMP1000/fields/orphans")
    assert orphan_page.status_code==200
    assert b"Orphan Supplier" in orphan_page.data
    assert b"250.0" in orphan_page.data

    restored=client.post("/panels/CMP1000/fields/orphans/rating/restore")
    assert restored.status_code==302
    restored_body=client.get("/api/panels/CMP1000").get_json()
    assert restored_body["panel"]["supplierFields"][0]["fieldId"]=="rating"
    assert "rating" not in restored_body["panel"]["metadata"]["orphanedSupplierFields"]

    client.post("/panels/CMP1000/configuration", data=edit_data)
    client.post("/panels/CMP1000/configuration", data={**edit_data,"confirm_field_removal":"1"})
    refused=client.post("/panels/CMP1000/fields/orphans/rating/purge", data={"confirm_field_id":"wrong"})
    assert refused.status_code==302
    assert "rating" in client.get("/api/panels/CMP1000").get_json()["panel"]["suppliers"][0]["customFields"]

    purged=client.post("/panels/CMP1000/fields/orphans/rating/purge", data={"confirm_field_id":"rating"})
    assert purged.status_code==302
    purged_body=client.get("/api/panels/CMP1000").get_json()
    assert "rating" not in purged_body["panel"]["suppliers"][0]["customFields"]
    assert "rating" not in purged_body["panel"]["metadata"]["orphanedSupplierFields"]
    assert purged_body["panel"]["metadata"]["auditTrail"][-1]["action"]=="orphan_field_purged"


def test_panel_import_migrates_v1_and_export_returns_current_schema(client):
    legacy={
        "schemaVersion":1,
        "panel":{
            "panelId":"CMP6000",
            "panelName":"Legacy Import",
            "category":"Transformers",
            "business":"GI",
            "region":{"level":"Global","value":"Global"},
            "panelOwner":{"name":"Owner"},
            "mdfCode":"MDF-TR-001",
            "supplierFields":[],
            "suppliers":[],
        },
    }
    response=client.post(
        "/panels/import",
        data={"panel_file":(io.BytesIO(json.dumps(legacy).encode("utf-8")),"legacy.json")},
        content_type="multipart/form-data",
    )
    assert response.status_code==302
    body=client.get("/api/panels/CMP6000").get_json()
    assert body["schemaVersion"]==5
    assert body["panel"]["leadMdfCode"]=="MDF-TR-001"
    assert body["panel"]["mdfCodes"]==["MDF-TR-001"]
    assert body["panel"]["metadata"]["version"]==5
    exported=client.get("/panels/CMP6000/export")
    assert exported.status_code==200
    assert exported.mimetype=="application/json"
    assert 'attachment; filename="CMP6000.json"' in exported.headers["Content-Disposition"]
    exported_body=json.loads(exported.data)
    assert exported_body["schemaVersion"]==5

def test_panel_import_requires_replace_for_existing_panel(client):
    create_panel(client)
    payload=client.get("/api/panels/CMP1000").get_json()
    payload["panel"]["panelName"]="Imported Replacement"
    blocked=client.post(
        "/panels/import",
        data={"panel_file":(io.BytesIO(json.dumps(payload).encode("utf-8")),"panel.json")},
        content_type="multipart/form-data",
    )
    assert blocked.status_code==400
    assert client.get("/api/panels/CMP1000").get_json()["panel"]["panelName"]=="Test Transformers"

    replaced=client.post(
        "/panels/import",
        data={
            "panel_file":(io.BytesIO(json.dumps(payload).encode("utf-8")),"panel.json"),
            "replace_existing":"1",
        },
        content_type="multipart/form-data",
    )
    assert replaced.status_code==302
    assert client.get("/api/panels/CMP1000").get_json()["panel"]["panelName"]=="Imported Replacement"

def test_panel_import_rejects_future_schema(client):
    payload={"schemaVersion":99,"panel":{}}
    response=client.post(
        "/panels/import",
        data={"panel_file":(io.BytesIO(json.dumps(payload).encode("utf-8")),"future.json")},
        content_type="multipart/form-data",
    )
    assert response.status_code==400
    assert b"Unsupported schemaVersion 99" in response.data


def test_seeded_demo_panels_use_historical_abb_europe_hub_data(tmp_path):
    seeded_app=create_app({
        "TESTING": True,
        "DATABASE": str(tmp_path / "seeded.sqlite"),
        "SEED_DEMO": True,
        "SECRET_KEY": "test",
    })
    seeded_client=seeded_app.test_client()

    portfolio=seeded_client.get("/")
    assert portfolio.status_code==200
    assert b"Power and Distribution Transformers" in portfolio.data
    assert b"LV &amp; MV Cables" in portfolio.data
    assert b"MV Switchgear" in portfolio.data
    assert b"Instrument Transformers" in portfolio.data
    assert b"Civil Works" in portfolio.data
    assert b"Engineering Services" in portfolio.data
    assert b"Protection &amp; Control Relays" in portfolio.data

    body=seeded_client.get("/api/panels/ABB-3GX").get_json()
    assert body["panel"]["region"]=={"level":"HUB","value":"Europe"}
    assert body["panel"]["metadata"]["legacySource"]=="Dynamic Supplier Panel 07Sep2021.xlsm"
    assert body["panel"]["metadata"]["legacySheet"]=="3GX"
    assert body["panel"]["leadMdfCode"]=="3GX"
    assert body["panel"]["suppliers"][0]["supplierName"]=="Koncar Power Transformers Ltd."
    assert body["panel"]["suppliers"][0]["address"]
    assert body["panel"]["suppliers"][0]["postCode"]
    assert body["panel"]["suppliers"][0]["customFields"]["country"]=="Croatia"
    assert body["panel"]["suppliers"][0]["customFields"]["classification"]=="2. HBU Preferred"
    assert body["panel"]["suppliers"][0]["customFields"]["ksm"]=="Emre Gul"
    assert body["panel"]["suppliers"][0]["customFields"]["supplier_risk"]=="Medium"

    civil=seeded_client.get("/api/panels/ABB-CIV").get_json()
    assert civil["panel"]["region"]=={"level":"HUB","value":"Europe"}
    field_ids={f["fieldId"] for f in civil["panel"]["supplierFields"]}
    assert {"branch_location","turnover_kusd","minimum_project_amount_kusd"} <= field_ids

    edit_page=seeded_client.get("/panels/ABB-3GX/edit")
    assert edit_page.status_code==200
    assert b"3GX" in edit_page.data


def test_audit_history_captures_panel_schema_supplier_and_raw_json_changes(client):
    created=create_panel(client)
    assert created.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    assert body["panel"]["metadata"]["auditTrail"][-1]["action"]=="panel_created"

    edited=client.post("/panels/CMP1000/edit", data={
        "panel_id":"CMP1000",
        "panel_name":"Test Transformers Updated",
        "category":"Transformers",
        "business":"GI",
        "region_level":"Country",
        "region_value":"United Kingdom",
        "owner":"New Owner",
        "mdf_codes":["MDF-TR-001"],
        "lead_mdf_code":"MDF-TR-001",
    })
    assert edited.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    update_event=body["panel"]["metadata"]["auditTrail"][-1]
    assert update_event["action"]=="panel_updated"
    assert "panelName" in update_event["details"]["changes"]
    assert "panelOwner" in update_event["details"]["changes"]
    assert "supplierFields" not in update_event["details"]["changes"]
    assert update_event["eventId"]
    assert update_event["actor"]=="local-user"

    configured=client.post("/panels/CMP1000/configuration", data={
        "field_groups_json":"[]",
        "fields_json":json.dumps([
            {"fieldId":"rating","fieldName":"Rating","type":"number","required":False,"options":[]},
            {"fieldId":"notes","fieldName":"Notes","type":"text","required":False,"options":[]},
        ]),
    })
    assert configured.status_code==302
    body=client.get("/api/panels/CMP1000").get_json()
    config_event=body["panel"]["metadata"]["auditTrail"][-1]
    assert config_event["action"]=="panel_configuration_updated"
    assert "supplierFields" in config_event["details"]["changes"]

    supplier=client.post("/panels/CMP1000/suppliers/new", data={
        "supplier_id":"SUP-AUDIT",
        "supplier_name":"Audit Supplier",
        "address":"Road",
        "post_code":"AA1 1AA",
        "custom_rating":"100",
        "custom_notes":"Created for audit",
    })
    assert supplier.status_code==302

    audit=client.get("/panels/CMP1000/audit")
    assert audit.status_code==200
    assert b"Panel Created" in audit.data
    assert b"Panel Updated" in audit.data
    assert b"Supplier Created" in audit.data
    assert b"Audit Supplier" in audit.data

    filtered=client.get("/panels/CMP1000/audit?action=supplier_created")
    assert filtered.status_code==200
    assert b"Supplier Created" in filtered.data
    assert filtered.data.count(b'class="audit-event"') == 1

    payload=client.get("/api/panels/CMP1000").get_json()
    prior_count=len(payload["panel"]["metadata"]["auditTrail"])
    payload["panel"]["panelName"]="Raw JSON Updated Name"
    raw=client.post("/panels/CMP1000/data", data={"raw_json":json.dumps(payload)})
    assert raw.status_code==302
    after=client.get("/api/panels/CMP1000").get_json()
    assert after["panel"]["panelName"]=="Raw JSON Updated Name"
    assert len(after["panel"]["metadata"]["auditTrail"])==prior_count+1
    assert after["panel"]["metadata"]["auditTrail"][-1]["action"]=="raw_json_updated"

def test_panel_import_records_audit_event(client):
    payload={
        "schemaVersion":3,
        "panel":{
            "panelId":"CMP-AUDIT-IMPORT",
            "panelName":"Imported Audit Panel",
            "category":"Transformers",
            "business":"GI",
            "region":{"level":"HUB","value":"Europe"},
            "panelOwner":{"name":"Europe Hub"},
            "mdfCode":"MDF-TR-001",
            "leadMdfCode":"MDF-TR-001",
            "mdfCodes":["MDF-TR-001"],
            "supplierFields":[],
            "suppliers":[],
            "metadata":{"version":3,"auditTrail":[]},
        },
    }
    response=client.post(
        "/panels/import",
        data={"panel_file":(io.BytesIO(json.dumps(payload).encode("utf-8")),"panel.json")},
        content_type="multipart/form-data",
    )
    assert response.status_code==302
    imported=client.get("/api/panels/CMP-AUDIT-IMPORT").get_json()
    assert imported["panel"]["metadata"]["auditTrail"][-1]["action"]=="panel_imported"


def test_portfolio_panel_qualification_uses_worst_supplier_status(client):
    fields=[
        {"fieldId":"qualification_status","fieldName":"Qualification Status","type":"text","required":False,"options":[]}
    ]
    response=client.post("/panels/new", data={
        "panel_id":"CMP-QUAL",
        "panel_name":"Qualification Demo",
        "category":"Transformers",
        "business":"GI",
        "region_level":"HUB",
        "region_value":"Europe",
        "owner":"Europe Hub",
        "mdf_codes":["MDF-TR-001"],
        "lead_mdf_code":"MDF-TR-001",
        "fields_json":json.dumps(fields),
    })
    assert response.status_code==302

    for sid,status in [
        ("SUP-Q1","Qualified"),
        ("SUP-Q2","In review"),
        ("SUP-Q3","Not qualified"),
    ]:
        added=client.post("/panels/CMP-QUAL/suppliers/new", data={
            "supplier_id":sid,
            "supplier_name":sid,
            "address":"Road",
            "post_code":"AA1 1AA",
            "custom_qualification_status":status,
        })
        assert added.status_code==302

    page=client.get("/?q=Qualification+Demo")
    assert page.status_code==200
    assert b"Not qualified" in page.data

def test_portfolio_panel_qualification_falls_back_to_in_review_then_qualified(client):
    fields=[
        {"fieldId":"qualification_status","fieldName":"Qualification Status","type":"text","required":False,"options":[]}
    ]
    client.post("/panels/new", data={
        "panel_id":"CMP-QUAL2",
        "panel_name":"Qualification Demo Two",
        "category":"Transformers",
        "business":"GI",
        "region_level":"HUB",
        "region_value":"Europe",
        "owner":"Europe Hub",
        "mdf_codes":["MDF-TR-001"],
        "lead_mdf_code":"MDF-TR-001",
        "fields_json":json.dumps(fields),
    })
    for sid,status in [("SUP-A","Qualified"),("SUP-B","In review")]:
        client.post("/panels/CMP-QUAL2/suppliers/new", data={
            "supplier_id":sid,
            "supplier_name":sid,
            "address":"Road",
            "post_code":"AA1 1AA",
            "custom_qualification_status":status,
        })
    page=client.get("/?q=Qualification+Demo+Two")
    assert b"In review" in page.data


def test_settings_hub_is_accessible_from_header(client):
    home=client.get("/")
    assert home.status_code==200
    assert b'href="/settings"' in home.data
    settings=client.get("/settings")
    assert settings.status_code==200
    assert b"Dashboard calculations" in settings.data
    assert b"Supplier master" in settings.data
    assert b"MDF master data" in settings.data
    assert b"Panel import" in settings.data

def test_duplicate_panel_preserves_configuration_but_not_suppliers_or_history(client):
    create_panel(client, "CMP-SOURCE")
    client.post("/panels/CMP-SOURCE/suppliers/new", data={
        "supplier_id": "SUP-DUP", "supplier_name": "Supplier to exclude",
        "address": "Test", "post_code": "", "custom_rating": "200",
    })
    source = client.get("/api/panels/CMP-SOURCE").get_json()["panel"]
    assert len(source["suppliers"]) == 1
    client.post("/panels/CMP-SOURCE/configuration", data={
        "field_groups_json": "[]",
        "fields_json": json.dumps(source["supplierFields"]),
        "dashboard_widgets_json": json.dumps([{"title": "Supplier Count", "metric": "count"}]),
    })
    source = client.get("/api/panels/CMP-SOURCE").get_json()["panel"]
    response = client.get("/panels/CMP-SOURCE/duplicate")
    assert response.status_code == 200
    assert b"Copy of Test Transformers" in response.data
    assert b"Duplicate Panel" in client.get("/panels/CMP-SOURCE").data

    response = client.post("/panels/CMP-SOURCE/duplicate", data={
        "panel_id": "CMP-COPY", "panel_name": "Independent copy",
        "category": "Transformers", "business": "GI", "region_level": "Country",
        "region_value": "United Kingdom", "owner": "Test Owner",
        "mdf_codes": ["MDF-TR-001"], "lead_mdf_code": "MDF-TR-001",
    })
    assert response.status_code == 302
    duplicate = client.get("/api/panels/CMP-COPY").get_json()["panel"]
    assert duplicate["suppliers"] == []
    assert duplicate["supplierFields"] == source["supplierFields"]
    assert duplicate["fieldGroups"] == source["fieldGroups"]
    assert duplicate["mdfCodes"] == source["mdfCodes"]
    assert duplicate["dashboardWidgets"] == source["dashboardWidgets"]
    assert len(duplicate["metadata"]["auditTrail"]) == 1
    assert duplicate["metadata"]["auditTrail"][0]["action"] == "panel_duplicated"
    assert duplicate["metadata"]["auditTrail"][0]["details"]["sourcePanelId"] == "CMP-SOURCE"
    assert len(client.get("/api/panels/CMP-SOURCE").get_json()["panel"]["suppliers"]) == 1
    assert client.post("/panels/CMP-SOURCE/duplicate", data={
        "panel_id": "CMP-COPY", "panel_name": "Duplicate", "category": "Transformers",
        "business": "GI", "region_level": "Country", "owner": "Test Owner",
        "mdf_codes": ["MDF-TR-001"], "lead_mdf_code": "MDF-TR-001",
    }).status_code == 409

def test_field_template_library_creates_independent_field_definitions(client):
    create_panel(client, "CMP-LIB")
    template = {"fieldName": "Cooling Type", "type": "dropdown",
                "options": ["ONAN", "ONAF"], "required": False}
    response = client.post("/api/field-templates", json=template)
    assert response.status_code == 201
    template_id = response.get_json()["templateId"]
    library = client.get("/api/field-templates").get_json()
    assert any(item["templateId"] == template_id for item in library)
    assert b"Select predefined field" in client.get("/panels/CMP-LIB/configuration").data
    assert client.post("/api/field-templates", json={"fieldName": "", "type": "number"}).status_code == 400

    fields = [{"fieldId": "cooling_copy", **template, "groupId": ""}]
    response = client.post("/panels/CMP-LIB/configuration", data={
        "field_groups_json": "[]", "fields_json": json.dumps(fields),
    })
    assert response.status_code == 302
    panel_field = client.get("/api/panels/CMP-LIB").get_json()["panel"]["supplierFields"][0]
    assert panel_field["fieldId"] == "cooling_copy"
    assert panel_field["options"] == ["ONAN", "ONAF"]
    # Updating a panel's field never modifies its central reusable template.
    fields[0]["options"].append("OFAF")
    client.post("/panels/CMP-LIB/configuration", data={
        "field_groups_json": "[]", "fields_json": json.dumps(fields),
    })
    assert next(item for item in client.get("/api/field-templates").get_json() if item["templateId"] == template_id)["options"] == ["ONAN", "ONAF"]

def test_boolean_field_tristate_in_supplier_and_bulk_editor(client):
    create_panel(client, "CMP-BOOL")
    client.post("/panels/CMP-BOOL/configuration", data={
        "field_groups_json": "[]",
        "fields_json": json.dumps([{
            "fieldId": "approved", "fieldName": "Approved", "type": "boolean",
            "required": False, "options": []
        }])
    })
    form = client.get("/panels/CMP-BOOL/suppliers/new")
    assert b'option value="true"' in form.data
    client.post("/panels/CMP-BOOL/suppliers/new", data={
        "supplier_id": "SUP-B", "supplier_name": "Test Boolean",
        "custom_approved": "true"
    })
    body = client.get("/api/panels/CMP-BOOL").get_json()
    assert body["panel"]["suppliers"][0]["customFields"]["approved"] is True
    assert b"boolean-yes" in client.get("/panels/CMP-BOOL").data
    client.post("/panels/CMP-BOOL/suppliers/bulk-edit", data={
        "SUP-B__approved": "false"
    })
    assert client.get("/api/panels/CMP-BOOL").get_json()["panel"]["suppliers"][0]["customFields"]["approved"] is False
    assert b"boolean-no" in client.get("/panels/CMP-BOOL").data
    client.post("/panels/CMP-BOOL/suppliers/SUP-B/edit", data={
        "supplier_name": "Test Boolean", "custom_approved": ""
    })
    assert client.get("/api/panels/CMP-BOOL").get_json()["panel"]["suppliers"][0]["customFields"]["approved"] is None
    assert b"boolean-unset" in client.get("/panels/CMP-BOOL").data

def test_configurable_supplier_kpi_widgets_and_safe_zero_division(client):
    create_panel(client, "CMP-KPI")
    fields = [
        {"fieldId": "mva", "fieldName": "MVA", "type": "number", "options": []},
        {"fieldId": "spend", "fieldName": "Spend", "type": "number", "options": []},
    ]
    widgets = [
        {"title": "Suppliers", "metric": "count"},
        {"title": "MVA", "metric": "sum", "fieldId": "mva"},
        {"title": "Total spend", "metric": "sum", "fieldId": "spend"},
        {"title": "Spend per MVA", "metric": "ratio",
         "fieldId": "spend", "otherFieldId": "mva"},
        {"title": "Average spend", "metric": "average", "fieldId": "spend"},
    ]
    saved = client.post("/panels/CMP-KPI/configuration", data={
        "fields_json": json.dumps(fields), "field_groups_json": "[]",
        "dashboard_widgets_json": json.dumps(widgets),
    })
    assert saved.status_code == 302
    client.post("/panels/CMP-KPI/suppliers/new", data={
        "supplier_id": "S1", "supplier_name": "One", "custom_mva": "100", "custom_spend": "500"
    })
    client.post("/panels/CMP-KPI/suppliers/new", data={
        "supplier_id": "S2", "supplier_name": "Two", "custom_mva": "200", "custom_spend": "900"
    })
    page = client.get("/panels/CMP-KPI")
    assert page.status_code == 200
    assert b"Spend per MVA" in page.data
    assert b"1,400.00" in page.data
    assert b"300.00" in page.data
    assert b"4.67" in page.data  # sum(spend)/sum(mva)
    assert b"700.00" in page.data
    saved_widgets = client.get("/api/panels/CMP-KPI").get_json()["panel"]["dashboardWidgets"]
    assert len(saved_widgets) == 5

    # Bad definitions rejected, and prior configuration preserved.
    failure = client.post("/panels/CMP-KPI/configuration", data={
        "fields_json": json.dumps(fields), "field_groups_json": "[]",
        "dashboard_widgets_json": json.dumps([{"title": f"W{i}", "metric": "count"} for i in range(7)]),
    })
    assert failure.status_code == 400
    assert len(client.get("/api/panels/CMP-KPI").get_json()["panel"]["dashboardWidgets"]) == 5
    client.post("/panels/CMP-KPI/configuration", data={
        "fields_json": json.dumps(fields), "field_groups_json": "[]",
        "dashboard_widgets_json": json.dumps([{
            "title": "Zero denominator", "metric": "ratio",
            "fieldId": "mva", "otherFieldId": "spend"
        }]),
    })
    assert client.get("/panels/CMP-KPI").status_code == 200

def test_currency_aware_kpi_sum_uses_base_currency_rates(client):
    create_panel(client, "CMP-FX")
    fields = [
        {"fieldId":"annual_spend","fieldName":"Spend","type":"number"},
        {"fieldId":"spend_currency","fieldName":"Currency","type":"dropdown",
         "options":["GBP","EUR"]},
    ]
    client.post("/panels/CMP-FX/configuration", data={
        "field_groups_json":"[]", "fields_json":json.dumps(fields),
        "dashboard_widgets_json":json.dumps([
            {"title":"Total spend","metric":"sum",
             "fieldId":"annual_spend","format":"currency"}
        ]),
    })
    for sid, currency in (("S1","GBP"),("S2","EUR")):
        client.post("/panels/CMP-FX/suppliers/new", data={
            "supplier_id":sid,"supplier_name":sid,
            "custom_annual_spend":"100","custom_spend_currency":currency,
        })
    response=client.get("/panels/CMP-FX")
    assert response.status_code==200
    assert "£186.00" in response.get_data(as_text=True)

def test_supplier_level_kpi_columns_and_display_modes(client):
    create_panel(client, "CMP-SUP-KPI")
    fields = [
        {"fieldId": "mva", "fieldName": "MVA", "type": "number"},
        {"fieldId": "annual_spend", "fieldName": "Spend", "type": "number"},
        {"fieldId": "spend_currency", "fieldName": "Currency", "type": "text"},
    ]
    widgets = [
        {"title": "Spend per MVA", "metric": "ratio", "fieldId": "annual_spend",
         "otherFieldId": "mva", "format": "currency", "display": "both"},
        {"title": "Supplier only", "metric": "sum", "fieldId": "mva",
         "display": "supplier"},
        {"title": "Panel only", "metric": "count", "display": "panel"},
    ]
    response = client.post("/panels/CMP-SUP-KPI/configuration", data={
        "field_groups_json": "[]", "fields_json": json.dumps(fields),
        "dashboard_widgets_json": json.dumps(widgets),
    })
    assert response.status_code == 302
    for supplier_id, mva, spend, currency in [
        ("S1", "100", "500", "GBP"),
        ("S2", "200", "900", "GBP"),
        ("S3", "0", "100", "GBP"),
        ("S4", "10", "100", "XYZ"),
    ]:
        assert client.post("/panels/CMP-SUP-KPI/suppliers/new", data={
            "supplier_id": supplier_id, "supplier_name": supplier_id,
            "custom_mva": mva, "custom_annual_spend": spend,
            "custom_spend_currency": currency,
        }).status_code == 302
    html = client.get("/panels/CMP-SUP-KPI").get_data(as_text=True)
    assert html.count('class="computed-kpi-head"') == 2
    assert html.count('class="computed-kpi-cell"') == 8
    assert "£5.00" in html
    assert "£4.50" in html
    assert 'class="empty-kpi"' in html
    assert "Panel only" in html
    assert "Supplier only" in html
    assert "£5.00" in html
    data = client.get("/api/panels/CMP-SUP-KPI").get_json()
    assert [x["display"] for x in data["panel"]["dashboardWidgets"]] == ["both", "supplier", "panel"]

    # Existing widget definitions without display remain panel-only.
    widgets[0].pop("display")
    response = client.post("/panels/CMP-SUP-KPI/configuration", data={
        "field_groups_json": "[]", "fields_json": json.dumps(fields),
        "dashboard_widgets_json": json.dumps(widgets),
    })
    assert response.status_code == 302
    html = client.get("/panels/CMP-SUP-KPI").get_data(as_text=True)
    assert html.count('class="computed-kpi-head"') == 1

def test_star_rating_supplier_edit_bulk_validation_and_kpi_display(client):
    create_panel(client, "CMP-STARS")
    fields = [{"fieldId":"quality","fieldName":"Quality","type":"stars","required":False,"options":[]}]
    widgets = [{"title":"Average Quality","metric":"average","fieldId":"quality","format":"stars","display":"both"}]
    assert client.post("/panels/CMP-STARS/configuration", data={
        "field_groups_json":"[]","fields_json":json.dumps(fields),
        "dashboard_widgets_json":json.dumps(widgets)
    }).status_code == 302
    assert b'value="stars"' in client.get("/panels/CMP-STARS/configuration").data or b"fields-json" in client.get("/panels/CMP-STARS/configuration").data
    assert client.post("/panels/CMP-STARS/suppliers/new",data={
        "supplier_id":"A","supplier_name":"Alpha","custom_quality":"4.5"
    }).status_code == 302
    assert client.post("/panels/CMP-STARS/suppliers/new",data={
        "supplier_id":"B","supplier_name":"Beta","custom_quality":"3.5"
    }).status_code == 302
    response = client.get("/panels/CMP-STARS")
    assert response.status_code == 200
    assert b"Average Quality" in response.data
    assert b"4.0 out of 5 stars" in response.data
    assert b"4.5 out of 5 stars" in response.data
    assert b"3.5 out of 5 stars" in response.data
    assert client.post("/panels/CMP-STARS/suppliers/A/edit",data={
        "supplier_name":"Alpha","custom_quality":"5.5"
    }).status_code == 400
    assert client.post("/panels/CMP-STARS/suppliers/bulk-edit",data={
        "A__quality":"4.3", "B__quality":"3"
    }).status_code == 400
    assert client.post("/panels/CMP-STARS/suppliers/bulk-edit",data={
        "A__quality":"5", "B__quality":"4"
    }).status_code == 302
    assert b"4.5 out of 5 stars" in client.get("/panels/CMP-STARS").data
    assert client.post("/panels/CMP-STARS/suppliers/new",data={
        "supplier_id":"C","supplier_name":"Invalid","custom_quality":"-1"
    }).status_code == 400
    assert any(x["type"] == "stars" for x in client.get("/api/field-templates").get_json())

def test_multiselect_blocks_one_field_many_coloured_segments(client):
    create_panel(client, "CMP-REGIONS")
    fields=[{"fieldId":"regional_coverage","fieldName":"Regional Coverage",
             "type":"multiselect_blocks","options":["EU","MED","MEA","NAM","LAM"]}]
    assert client.post("/panels/CMP-REGIONS/configuration",data={
        "field_groups_json":"[]","fields_json":json.dumps(fields)
    }).status_code==302
    page=client.get("/panels/CMP-REGIONS/suppliers/new")
    assert b'value="EU"' in page.data and b'value="LAM"' in page.data
    assert client.post("/panels/CMP-REGIONS/suppliers/new",data={
        "supplier_id":"S1","supplier_name":"One",
        "custom_regional_coverage":["EU","MED","LAM"]
    }).status_code==302
    payload=client.get("/api/panels/CMP-REGIONS").get_json()["panel"]
    assert payload["suppliers"][0]["customFields"]["regional_coverage"]==["EU","MED","LAM"]
    html=client.get("/panels/CMP-REGIONS").get_data(as_text=True)
    assert html.count("multi-block-segment")==5
    assert html.count("multi-block-active")==3
    assert html.count("multi-block-inactive")==2
    assert client.post("/panels/CMP-REGIONS/suppliers/bulk-edit",data={
        "S1__regional_coverage":["MEA","NAM"]
    }).status_code==302
    assert client.get("/api/panels/CMP-REGIONS").get_json()["panel"]["suppliers"][0]["customFields"]["regional_coverage"]==["MEA","NAM"]
    assert client.post("/panels/CMP-REGIONS/suppliers/S1/edit",data={
        "supplier_name":"One","custom_regional_coverage":["UNKNOWN"]
    }).status_code==400
    assert client.get("/api/panels/CMP-REGIONS").get_json()["panel"]["suppliers"][0]["customFields"]["regional_coverage"]==["MEA","NAM"]
    assert client.post("/panels/CMP-REGIONS/suppliers/S1/edit",data={"supplier_name":"One"}).status_code==302
    assert client.get("/api/panels/CMP-REGIONS").get_json()["panel"]["suppliers"][0]["customFields"]["regional_coverage"]==[]

def test_star_kpi_selects_custom_field_aggregation_and_output(client):
    create_panel(client, "CMP-RATING-KPI")
    fields=[{"fieldId":"quality","fieldName":"Quality Rating","type":"stars"},
            {"fieldId":"score","fieldName":"Score","type":"number"},
            {"fieldId":"notes","fieldName":"Notes","type":"text"}]
    widgets=[
        {"title":"Average Quality","metric":"average","fieldId":"quality","format":"stars_both","display":"both"},
        {"title":"Worst Quality","metric":"minimum","fieldId":"quality","format":"stars","display":"panel"},
        {"title":"Best Quality","metric":"maximum","fieldId":"quality","format":"number","display":"panel"}
    ]
    def save(widgets):
        return client.post("/panels/CMP-RATING-KPI/configuration",data={
            "field_groups_json":"[]","fields_json":json.dumps(fields),
            "dashboard_widgets_json":json.dumps(widgets)
        })
    assert save(widgets).status_code==302
    assert b"stars_both" in client.get("/panels/CMP-RATING-KPI/configuration").data
    for sid,rating in [("A","3"),("B","4.5"),("C","5")]:
        assert client.post("/panels/CMP-RATING-KPI/suppliers/new",data={
            "supplier_id":sid,"supplier_name":sid,"custom_quality":rating
        }).status_code==302
    html=client.get("/panels/CMP-RATING-KPI").get_data(as_text=True)
    assert "Average Quality" in html and "4.2 out of 5 stars" in html
    assert "Worst Quality" in html and "3.0 out of 5 stars" in html
    assert "Best Quality" in html and "5.00" in html
    widgets[0]["fieldId"]="notes"
    assert save(widgets).status_code==400
    widgets[0]["fieldId"]="quality"
    widgets[0]["metric"]="ratio"
    widgets[0]["otherFieldId"]="score"
    assert save(widgets).status_code==400

def test_mdf_search_rows_can_be_hidden_by_client_filter():
    """The picker hides non-matching rows via hidden; CSS must not force them visible."""
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    css = (root / "static" / "app.css").read_text()
    js = (root / "static" / "app.js").read_text()
    template = (root / "templates" / "panel_form.html").read_text()
    assert ".mdf-option[hidden]{display:none!important;}" in css
    assert "row.hidden=q && !row.dataset.search.includes(q)" in js
    assert 'data-search="{{ (m.code ~ \' \' ~ m.description)|lower }}"' in template

def test_supplier_table_filter_sort_controls_for_custom_types(client):
    create_panel(client, "CMP-FILTER")
    fields = [
        {"fieldId": "rating", "fieldName": "Quality", "type": "stars"},
        {"fieldId": "coverage", "fieldName": "Regions", "type": "multiselect_blocks",
         "options": ["EU", "NAM", "MEA"]},
        {"fieldId": "approved", "fieldName": "Approved", "type": "boolean"},
    ]
    response = client.post("/panels/CMP-FILTER/configuration", data={
        "field_groups_json": "[]", "fields_json": json.dumps(fields),
        "dashboard_widgets_json": "[]",
    })
    assert response.status_code == 302
    assert client.post("/panels/CMP-FILTER/suppliers/new", data={
        "supplier_id": "SUP1", "supplier_name": "Alpha",
        "custom_rating": "4.5", "custom_coverage": ["EU", "MEA"], "custom_approved": "true",
    }).status_code == 302
    page = client.get("/panels/CMP-FILTER")
    html = page.get_data(as_text=True)
    assert page.status_code == 200
    assert 'data-supplier-controls' in html
    assert 'data-supplier-table' in html
    assert 'class="supplier-table-search"' in html
    assert 'class="supplier-filter-field"' in html
    assert 'class="supplier-sort-field"' in html
    assert 'data-type="number"' in html
    assert 'data-filter="EU MEA"' in html
    assert 'data-filter="Yes"' in html
    # UI filtering must never remove supplier information from server-side storage.
    assert len(client.get("/api/panels/CMP-FILTER").get_json()["panel"]["suppliers"]) == 1


def test_supplier_filter_visibility_and_typed_sort_js_contract():
    from pathlib import Path
    root = Path(__file__).resolve().parents[1]
    js = (root / "static" / "app.js").read_text()
    css = (root / "static" / "app.css").read_text()
    assert 'normalise(cell.dataset.filter??cell.textContent).includes(needle)' in js
    assert 'ca.dataset.type==="number"' in js
    assert 'row.hidden=!match' in js
    assert '.supplier-data-row[hidden]{display:none!important;}' in css

def test_dev033_supplier_column_picker_preserves_supplier_actions(client):
    create_panel(client, "CMP-COLUMNS")
    page=client.get("/panels/CMP-COLUMNS")
    assert page.status_code==200
    html=page.get_data(as_text=True)
    assert 'data-supplier-column-settings' in html
    assert 'data-panel-id="CMP-COLUMNS"' in html
    assert 'class="supplier-column-options"' in html
    assert 'supplier-column-reset' in html
    assert 'name="confirm_supplier_id"' in html or 'class="supplier-data-row"' not in html


def test_dev033_column_order_keeps_original_filter_and_sort_indices():
    from pathlib import Path
    root=Path(__file__).resolve().parents[1]
    js=(root/"static"/"app.js").read_text()
    css=(root/"static"/"app.css").read_text()
    assert 'const rowCells=new Map(rows.map(row=>[row,Array.from(row.cells)]))' in js
    assert 'const ca=originalCell(a,sortIndex),cb=originalCell(b,sortIndex)' in js
    assert 'const cells=rowCells.get(row).slice(1,-1)' in js
    assert 'localStorage.setItem(storageKey' in js
    assert 'columnOrder=[0,1,2,...order.filter(i=>!fixed.has(i)),headings.length-1]' in js
    assert '[data-supplier-table] th[hidden],[data-supplier-table] td[hidden]{display:none!important;}' in css

def test_dev037_unsaved_guard_on_editing_forms(client):
    create_panel(client,"CMP-UNSAVED")
    urls=[
        "/panels/new",
        "/panels/CMP-UNSAVED/edit",
        "/panels/CMP-UNSAVED/configuration",
        "/panels/CMP-UNSAVED/suppliers/new",
        "/panels/CMP-UNSAVED/suppliers/bulk-edit",
    ]
    for url in urls:
        response=client.get(url)
        assert response.status_code==200
        assert "data-unsaved-guard" in response.get_data(as_text=True)


def test_dev037_dirty_navigation_guard_js_contract():
    from pathlib import Path
    js=(Path(__file__).resolve().parents[1]/"static"/"app.js").read_text()
    assert 'form.addEventListener("input",markDirty)' in js
    assert 'form.addEventListener("change",markDirty)' in js
    assert 'window.addEventListener("beforeunload"' in js
    assert 'if(dirty&&!submitting)' in js
    assert 'form.addEventListener("submit",()=>{submitting=true;})' in js
    assert 'Leave without saving?' in js

def test_supplier_excel_export_preview_confirm_and_conflict(client):
    import io
    from openpyxl import load_workbook
    create_panel(client,"CMP-EXCEL")
    assert client.post("/panels/CMP-EXCEL/suppliers/new",data={
        "supplier_id":"BP1","supplier_name":"Supplier One","custom_rating":"3"
    }).status_code==302
    export=client.get("/panels/CMP-EXCEL/suppliers/excel")
    assert export.status_code==200
    assert export.data.startswith(b"PK")
    book=load_workbook(io.BytesIO(export.data))
    ws=book["Suppliers"]
    assert ws["A1"].value=="Supplier ID"
    assert ws["C1"].value=="rating"
    ws["C2"]=4.5
    updated=io.BytesIO()
    book.save(updated)
    preview=client.post("/panels/CMP-EXCEL/suppliers/excel",data={
        "workbook":(io.BytesIO(updated.getvalue()),"changes.xlsx")
    },content_type="multipart/form-data")
    assert preview.status_code==200
    assert b"Review supplier Excel changes" in preview.data
    assert b"BP1" in preview.data
    assert client.get("/api/panels/CMP-EXCEL").get_json()["panel"]["suppliers"][0]["customFields"]["rating"]==3.0
    import re
    token=re.search(rb'name="token" value="([^"]+)"',preview.data).group(1).decode()
    commit=client.post("/panels/CMP-EXCEL/suppliers/excel/confirm",data={"token":token})
    assert commit.status_code==302
    assert client.get("/api/panels/CMP-EXCEL").get_json()["panel"]["suppliers"][0]["customFields"]["rating"]==4.5
    assert "supplier_excel_import" in [x["action"] for x in client.get("/api/panels/CMP-EXCEL").get_json()["panel"]["metadata"]["auditTrail"]]
    # Previous export baseline is now stale; it must not overwrite a newer edit.
    old=client.post("/panels/CMP-EXCEL/suppliers/excel",data={
        "workbook":(io.BytesIO(updated.getvalue()),"old.xlsx")
    },content_type="multipart/form-data",follow_redirects=True)
    assert b"changed since workbook export" in old.data


def test_supplier_excel_rejects_schema_mismatch_and_invalid_rating(client):
    import io
    from openpyxl import load_workbook
    create_panel(client,"CMP-EXCEL-ERR")
    assert client.post("/panels/CMP-EXCEL-ERR/suppliers/new",data={
        "supplier_id":"BP1","supplier_name":"One","custom_rating":"2"
    }).status_code==302
    raw=client.get("/panels/CMP-EXCEL-ERR/suppliers/excel").data
    book=load_workbook(io.BytesIO(raw))
    ws=book["Suppliers"]
    ws["C1"]="bad_field"
    invalid=io.BytesIO();book.save(invalid)
    response=client.post("/panels/CMP-EXCEL-ERR/suppliers/excel",data={
        "workbook":(io.BytesIO(invalid.getvalue()),"invalid.xlsx")
    },content_type="multipart/form-data",follow_redirects=True)
    assert b"Column headers" in response.data
    assert client.get("/api/panels/CMP-EXCEL-ERR").get_json()["panel"]["suppliers"][0]["customFields"]["rating"]==2.0

def test_dev035_live_kpi_preview_without_saving(client):
    create_panel(client, "CMP-PREVIEW")
    assert client.post("/panels/CMP-PREVIEW/suppliers/new",data={
        "supplier_id":"S-1","supplier_name":"One","custom_rating":"4"
    }).status_code==302
    assert client.post("/panels/CMP-PREVIEW/suppliers/new",data={
        "supplier_id":"S-2","supplier_name":"Two","custom_rating":"2"
    }).status_code==302
    widget={"title":"Average rating","metric":"average","fieldId":"rating",
            "format":"number","display":"panel"}
    preview=client.post("/panels/CMP-PREVIEW/widgets/preview",json={"widgets":[widget]})
    assert preview.status_code==200
    assert preview.get_json()["widgets"][0]["value"]==3.0
    panel=client.get("/api/panels/CMP-PREVIEW").get_json()["panel"]
    assert panel.get("dashboardWidgets",[])==[]
    invalid=client.post("/panels/CMP-PREVIEW/widgets/preview",json={
        "widgets":[dict(widget,fieldId="missing")]
    })
    assert invalid.status_code==400
    assert "unknown supplier field" in invalid.get_json()["error"]
    page=client.get("/panels/CMP-PREVIEW/configuration")
    assert b'data-preview-url=' in page.data
    assert b"current supplier data" in page.data

def test_dev036_supplier_comparison_is_read_only_and_validated(client):
    create_panel(client, "CMP-COMPARE")
    for sid, rating in [("ABC","3"),("XYZ","4")]:
        assert client.post("/panels/CMP-COMPARE/suppliers/new",data={
            "supplier_id":sid,"supplier_name":sid,"custom_rating":rating
        }).status_code==302
    page=client.get("/panels/CMP-COMPARE")
    assert page.status_code==200
    assert page.data.count(b'class="supplier-compare-check"')==2
    assert b'supplier-compare-form' in page.data
    before=client.get("/api/panels/CMP-COMPARE").get_json()
    response=client.get("/panels/CMP-COMPARE/suppliers/compare?supplier_id=ABC&supplier_id=XYZ")
    assert response.status_code==200
    assert b"Supplier comparison" in response.data
    assert b"Rating" in response.data
    assert b"ABC" in response.data and b"XYZ" in response.data
    assert client.get("/api/panels/CMP-COMPARE").get_json()==before
    assert client.get("/panels/CMP-COMPARE/suppliers/compare?supplier_id=ABC").status_code==302
    assert client.get("/panels/CMP-COMPARE/suppliers/compare?supplier_id=ABC&supplier_id=ABC").status_code==302
    assert client.get("/panels/CMP-COMPARE/suppliers/compare?supplier_id=ABC&supplier_id=UNKNOWN").status_code==400


def test_dev036_selection_preserves_table_filters_and_column_layout():
    from pathlib import Path
    root=Path(__file__).resolve().parents[1]
    js=(root/"static"/"app.js").read_text()
    assert 'const fixed=new Set([0,1,2,headings.length-1]);' in js
    assert 'const candidates=filterIndex===null?cells:[originalCell(row,filterIndex)];' in js
    assert 'toggles.forEach(box=>box.addEventListener("change",refreshComparison))' in js

def test_dev038_weighted_supplier_scoring_configuration_and_missing_values(client):
    from supplier_scoring import validate_scoring, score_suppliers
    fields=[{"fieldId":"quality","fieldName":"Quality","type":"stars"},
            {"fieldId":"delivery","fieldName":"Delivery","type":"number"},
            {"fieldId":"region","fieldName":"Region","type":"text"}]
    rules=validate_scoring([{"fieldId":"quality","weight":60},{"fieldId":"delivery","weight":40}],fields)
    core={"supplierFields":fields,"scoringCriteria":rules,"suppliers":[
        {"supplierId":"A","customFields":{"quality":5,"delivery":2.5}},
        {"supplierId":"B","customFields":{"quality":4}},
    ]}
    rows=score_suppliers(core)
    assert rows[0]["score"]==4.0
    assert rows[1]["score"] is None and not rows[1]["complete"]
    for invalid in [[{"fieldId":"region","weight":2}],
                    [{"fieldId":"quality","weight":0}],
                    [{"fieldId":"quality","weight":1},{"fieldId":"quality","weight":2}]]:
        with pytest.raises(ValueError):
            validate_scoring(invalid,fields)


def test_dev038_panel_scoring_saved_and_displayed(client):
    create_panel(client,"CMP-SCORE")
    fields=[{"fieldId":"rating","fieldName":"Rating","type":"number","required":False,"options":[]}]
    response=client.post("/panels/CMP-SCORE/configuration",data={
        "field_groups_json":"[]","fields_json":json.dumps(fields),
        "dashboard_widgets_json":"[]",
        "scoring_criteria_json":json.dumps([{"fieldId":"rating","weight":100}])
    })
    assert response.status_code==302
    core=client.get("/api/panels/CMP-SCORE").get_json()["panel"]
    assert core["scoringCriteria"]==[{"fieldId":"rating","weight":100.0}]
    assert client.post("/panels/CMP-SCORE/suppliers/new",data={
        "supplier_id":"S1","supplier_name":"Supplier One","custom_rating":"4.5"
    }).status_code==302
    page=client.get("/panels/CMP-SCORE?tab=scoring")
    assert b"Weighted supplier scores" in page.data
    assert b"4.5 / 5" in page.data
    invalid=client.post("/panels/CMP-SCORE/configuration",data={
        "field_groups_json":"[]","fields_json":json.dumps(fields),
        "dashboard_widgets_json":"[]",
        "scoring_criteria_json":json.dumps([{"fieldId":"rating","weight":-5}])
    })
    assert invalid.status_code==400
    assert client.get("/api/panels/CMP-SCORE").get_json()["panel"]["scoringCriteria"]==core["scoringCriteria"]

def test_dev039_supplier_actions_create_update_audit_and_validation(client):
    create_panel(client, "CMP-ACTIONS")
    assert client.post("/panels/CMP-ACTIONS/suppliers/new", data={
        "supplier_id": "SUP-A", "supplier_name": "Action Supplier"
    }).status_code == 302
    url = "/panels/CMP-ACTIONS/suppliers/SUP-A/actions"
    assert b"No actions recorded." in client.get(url).data
    created = client.post(url, data={
        "title": "Obtain quality certificate", "owner": "Category Manager",
        "due_date": "2026-12-15", "status": "Open"
    })
    assert created.status_code == 302
    core = client.get("/api/panels/CMP-ACTIONS").get_json()["panel"]
    action = core["suppliers"][0]["actions"][0]
    assert action["title"] == "Obtain quality certificate"
    assert action["owner"] == "Category Manager"
    assert action["dueDate"] == "2026-12-15"
    assert action["status"] == "Open"
    assert b"Actions (1)" in client.get("/panels/CMP-ACTIONS").data
    assert b"Obtain quality certificate" in client.get(url).data
    updated = client.post(url, data={
        "action_id": action["actionId"], "title": "Certificate received",
        "owner": "Quality Lead", "due_date": "2026-12-15",
        "status": "Completed"
    })
    assert updated.status_code == 302
    core = client.get("/api/panels/CMP-ACTIONS").get_json()["panel"]
    actions = core["suppliers"][0]["actions"]
    assert len(actions) == 1
    assert actions[0]["status"] == "Completed"
    assert actions[0]["title"] == "Certificate received"
    events = [event["action"] for event in core["metadata"]["auditTrail"]]
    assert "supplier_action_created" in events
    assert "supplier_action_updated" in events
    bad = client.post(url, data={
        "title": "Invalid", "owner": "Test", "due_date": "2026-13-75", "status": "Open"
    })
    assert bad.status_code == 400
    assert len(client.get("/api/panels/CMP-ACTIONS").get_json()["panel"]["suppliers"][0]["actions"]) == 1
    assert client.post(url, data={
        "action_id": "unknown", "title": "Nope", "owner": "Test",
        "status": "Open"
    }).status_code == 404
    assert client.get("/panels/CMP-ACTIONS/suppliers/missing/actions").status_code == 404


def test_dev039_supplier_action_rejects_unsafe_status_and_empty_owner(client):
    create_panel(client, "CMP-ACT-ERR")
    client.post("/panels/CMP-ACT-ERR/suppliers/new", data={
        "supplier_id": "S1", "supplier_name": "Supplier"
    })
    url = "/panels/CMP-ACT-ERR/suppliers/S1/actions"
    for changes in ({"title": "Task", "owner": "", "status": "Open"},
                    {"title": "Task", "owner": "Lead", "status": "Escalated"}):
        assert client.post(url, data=changes).status_code == 400
    assert client.get("/api/panels/CMP-ACT-ERR").get_json()["panel"]["suppliers"][0].get("actions", []) == []

def test_dev040_review_alert_classification_and_actions():
    from datetime import date
    from supplier_alerts import classify_due_date, supplier_alerts
    today=date(2026,10,8)
    assert classify_due_date("2026-10-07",today)=="Overdue"
    assert classify_due_date("2026-10-08",today)=="Due soon"
    assert classify_due_date("2026-11-07",today)=="Due soon"
    assert classify_due_date("2026-11-08",today)=="Future"
    assert classify_due_date("",today)=="Missing"
    assert classify_due_date("not-a-date",today)=="Invalid"
    core={"suppliers":[{"supplierId":"A","supplierName":"Alpha",
        "customFields":{"review":"2026-10-07"},
        "actions":[{"title":"Renew","status":"Open","dueDate":"2026-10-06"},
                   {"title":"Done","status":"Completed","dueDate":"2026-10-06"}]}]}
    rows=supplier_alerts(core,"review",today)
    assert rows[0]["reviewStatus"]=="Overdue"
    assert len(rows[0]["actions"])==1
    assert rows[0]["actions"][0]["status"]=="Overdue"


def test_dev040_panel_renders_qualification_alerts(client):
    create_panel(client,"CMP-ALERT")
    response=client.get("/panels/CMP-ALERT?tab=actions")
    assert response.status_code==200
    assert b"Qualification review alerts" in response.data
    assert b"No qualification reviews or actions require attention." in response.data
    assert b"Review status" in response.data

def test_dev041_coverage_concentration_and_qualification():
    from supplier_coverage import coverage_summary
    core={"mdfCodes":["MDF-TR-001"],"leadMdfCode":"MDF-TR-001",
          "supplierFields":[{"fieldId":"regions","fieldName":"Regional coverage",
                             "type":"multiselect_blocks","options":["EU","MEA","APAC"]}],
          "suppliers":[
            {"supplierId":"A","customFields":{"regions":["EU","MEA"],"qualification_status":"Qualified"}},
            {"supplierId":"B","customFields":{"regions":["EU"],"qualification_status":"In review"}},
            {"supplierId":"C","customFields":{"regions":"APAC","qualification_status":"Qualified"}},
          ]}
    result=coverage_summary(core)
    regions={r["region"]:r for r in result["groups"][0]["regions"]}
    assert (regions["EU"]["count"],regions["EU"]["qualifiedCount"],regions["EU"]["risk"])==(2,1,"Multiple sources")
    assert (regions["MEA"]["count"],regions["MEA"]["qualifiedCount"],regions["MEA"]["risk"])==(1,1,"Single source")
    assert regions["APAC"]["risk"]=="Uncovered"
    assert result["mdfCodes"]==["MDF-TR-001"]
    assert all("supplierMdf" not in row for row in regions.values())


def test_dev041_panel_coverage_renders_without_configured_regions(client):
    create_panel(client,"CMP-COVERAGE")
    response=client.get("/panels/CMP-COVERAGE?tab=overview")
    assert response.status_code==200
    assert b"Supplier coverage and sourcing concentration" in response.data
    assert b"No multi-select coverage field is configured." in response.data

def test_dev042_full_panel_template_roundtrip_without_supplier_data(client):
    create_panel(client,"CMP-TEMPLATE-SOURCE")
    assert client.post("/panels/CMP-TEMPLATE-SOURCE/suppliers/new", data={
        "supplier_id":"S1","supplier_name":"Original supplier","custom_rating":"4"
    }).status_code==302
    assert client.post("/panels/CMP-TEMPLATE-SOURCE/configuration",data={
        "field_groups_json":"[]",
        "fields_json":json.dumps([{"fieldId":"rating","fieldName":"Rating","type":"number","required":False,"options":[]}]),
        "dashboard_widgets_json":json.dumps([{"title":"Average rating","metric":"average","fieldId":"rating","format":"number","display":"panel"}]),
        "scoring_criteria_json":json.dumps([{"fieldId":"rating","weight":100}])
    }).status_code==302
    resp=client.post("/panels/CMP-TEMPLATE-SOURCE/save-template",data={"name":"Standard rating panel"})
    assert resp.status_code==302
    import sqlite3
    with sqlite3.connect(client.application.config["DATABASE"]) as conn:
        conn.row_factory=sqlite3.Row
        templates=conn.execute("SELECT template_id,config_json FROM panel_templates").fetchall()
        assert len(templates)==1
        tid=templates[0]["template_id"]
        saved=json.loads(templates[0]["config_json"])
        assert "suppliers" not in saved and "metadata" not in saved
    page=client.get("/panel-templates")
    assert b"Standard rating panel" in page.data
    form=client.get("/panel-templates/"+tid+"/create")
    assert form.status_code==200
    created=client.post("/panel-templates/"+tid+"/create",data={
        "panel_id":"CMP-TEMPLATE-NEW","panel_name":"New copy",
        "category":"Transformers","business":"GI","region_level":"Country",
        "region_value":"United Kingdom","owner":"New Owner",
        "mdf_codes":["MDF-TR-001"],"lead_mdf_code":"MDF-TR-001",
    })
    assert created.status_code==302
    original=client.get("/api/panels/CMP-TEMPLATE-SOURCE").get_json()["panel"]
    duplicate=client.get("/api/panels/CMP-TEMPLATE-NEW").get_json()["panel"]
    assert duplicate["supplierFields"]==original["supplierFields"]
    assert duplicate["dashboardWidgets"]==original["dashboardWidgets"]
    assert duplicate["scoringCriteria"]==original["scoringCriteria"]
    assert duplicate["suppliers"]==[]
    assert "Original supplier" not in client.get("/panels/CMP-TEMPLATE-NEW").get_data(as_text=True)
    assert client.post("/panels/CMP-TEMPLATE-SOURCE/save-template",data={"name":""}).status_code==302
    assert client.get("/panel-templates/not-a-template/create").status_code==404

def test_dev043_supplier_history_audit_sources_and_numeric_trend():
    from supplier_history import supplier_history
    core={"supplierFields":[{"fieldId":"rating","fieldName":"Rating","type":"stars"}],
          "metadata":{"auditTrail":[
              {"timestamp":"2026-10-01T12:00:00","actor":"tester","action":"supplier_updated","entityId":"A",
               "details":{"before":{"customFields":{"rating":2}},"after":{"customFields":{"rating":4}}}},
              {"timestamp":"2026-10-02T12:00:00","actor":"tester","action":"supplier_excel_import","entityId":"PANEL",
               "details":{"changes":[{"supplierId":"A","before":{"rating":4},"after":{"rating":5}},
                                     {"supplierId":"B","before":{"rating":1},"after":{"rating":3}}]}},
              {"timestamp":"2026-10-03T12:00:00","actor":"tester","action":"supplier_action_updated","entityId":"A",
               "details":{"before":{"actionId":"x","status":"Open"},"after":{"actionId":"x","status":"Completed"}}},
          ]}}
    history=supplier_history(core,"A")
    assert len(history["events"])==3
    assert [(p["value"]) for p in history["trends"][0]["points"]]==[4.0,5.0]
    assert history["events"][0]["changes"][0]["field"]=="action.status"
    assert history["events"][1]["changes"][0]["label"]=="Rating"
    assert len(supplier_history(core,"B")["events"])==1


def test_dev043_supplier_history_page_read_only(client):
    create_panel(client,"CMP-HISTORY")
    client.post("/panels/CMP-HISTORY/suppliers/new",data={"supplier_id":"S1","supplier_name":"One","custom_rating":"2"})
    client.post("/panels/CMP-HISTORY/suppliers/S1/edit",data={
        "supplier_name":"One","address":"New address","post_code":"ST1","custom_rating":"4"})
    before=client.get("/api/panels/CMP-HISTORY").get_json()
    page=client.get("/panels/CMP-HISTORY/suppliers/S1/history")
    assert page.status_code==200
    assert b"Supplier history and trends" in page.data
    assert b"Rating" in page.data
    assert b"History" in client.get("/panels/CMP-HISTORY").data
    assert client.get("/api/panels/CMP-HISTORY").get_json()==before
    assert client.get("/panels/CMP-HISTORY/suppliers/UNKNOWN/history").status_code==404

def test_dev044_management_excel_report_without_writing_panel(client):
    import io
    from openpyxl import load_workbook
    create_panel(client, "CMP-REPORT")
    assert client.post("/panels/CMP-REPORT/suppliers/new",data={
        "supplier_id":"ABC","supplier_name":"Sample Supplier","custom_rating":"4"
    }).status_code==302
    before=client.get("/api/panels/CMP-REPORT").get_json()
    response=client.get("/panels/CMP-REPORT/management-report.xlsx")
    assert response.status_code==200
    assert response.data.startswith(b"PK")
    assert "attachment" in response.headers["Content-Disposition"]
    workbook=load_workbook(io.BytesIO(response.data),read_only=True,data_only=True)
    assert workbook.sheetnames==["Overview","Suppliers","Regional coverage","Supplier actions"]
    assert workbook["Overview"]["B3"].value=="CMP-REPORT"
    assert workbook["Overview"]["B6"].value==1
    assert workbook["Suppliers"]["A2"].value=="ABC"
    assert workbook["Suppliers"]["B2"].value=="Sample Supplier"
    assert workbook["Suppliers"]["E2"].value=="Incomplete / not configured"
    assert client.get("/api/panels/CMP-REPORT").get_json()==before
    assert client.get("/panels/UNKNOWN/management-report.xlsx").status_code==404


def test_dev044_report_includes_qualified_coverage_and_open_actions():
    import io
    from openpyxl import load_workbook
    from supplier_coverage import coverage_summary
    from supplier_alerts import supplier_alerts
    from supplier_scoring import score_suppliers
    from management_report import build_management_workbook
    core={"panelId":"TEST","panelName":"Test","mdfCodes":["MDF-01"],
        "leadMdfCode":"MDF-01",
        "supplierFields":[{"fieldId":"regions","fieldName":"Regions","type":"multiselect_blocks","options":["EU","MEA"]}],
        "suppliers":[{"supplierId":"S1","supplierName":"Supplier",
           "customFields":{"regions":["EU"],"qualification_status":"Qualified","qualification_review_date":"2026-10-01"},
           "actions":[{"title":"Obtain certification","owner":"Manager","dueDate":"2026-11-01","status":"Open"},
                      {"title":"Old action","owner":"Manager","status":"Completed"}]}]}
    content=build_management_workbook(core,score_suppliers(core),
           coverage_summary(core),supplier_alerts(core,"qualification_review_date"),
           "qualification_status")
    book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    assert book["Regional coverage"]["D2"].value==1
    assert book["Regional coverage"]["E3"].value=="Uncovered"
    assert book["Supplier actions"]["C2"].value=="Obtain certification"
    assert book["Supplier actions"]["C3"].value is None

def test_dev044_printable_management_summary(client):
    create_panel(client,"CMP-PRINT")
    response=client.get("/panels/CMP-PRINT/management-report")
    assert response.status_code==200
    assert b"Supplier Management Summary" in response.data
    assert b"Print / Save as PDF" in response.data
    assert b"Regional sourcing coverage" in response.data
    assert b"Outstanding supplier actions" in response.data
    assert b"Print Management Report" in client.get("/panels/CMP-PRINT").data

def test_dev045_supplier_profile_joins_exact_id_across_panels(client):
    create_panel(client, "CMP-PROFILE-A")
    create_panel(client, "CMP-PROFILE-B")
    create_panel(client, "CMP-PROFILE-C")
    for pid, sid, name in [
        ("CMP-PROFILE-A", "BP100", "Example Ltd"),
        ("CMP-PROFILE-B", "BP100", "Example Trading"),
        ("CMP-PROFILE-C", "BP200", "Example Ltd"),
    ]:
        assert client.post(f"/panels/{pid}/suppliers/new",data={
            "supplier_id":sid,"supplier_name":name,"custom_rating":"3"
        }).status_code==302
    before_a=client.get("/api/panels/CMP-PROFILE-A").get_json()
    response=client.get("/suppliers/BP100/profile")
    assert response.status_code==200
    page=response.get_data(as_text=True)
    assert "Cross-panel supplier profile" in page
    assert "CMP-PROFILE-A" in page and "CMP-PROFILE-B" in page
    assert "CMP-PROFILE-C" not in page
    assert "Example Trading" in page
    assert "Profile" in client.get("/panels/CMP-PROFILE-A").get_data(as_text=True)
    assert client.get("/api/panels/CMP-PROFILE-A").get_json()==before_a
    assert client.get("/suppliers/UNKNOWN/profile").status_code==404

def test_dev046_quality_checks_find_missing_invalid_orphan_fields():
    from data_quality import inspect_panel
    core={"supplierFields":[
        {"fieldId":"rating","fieldName":"Rating","type":"stars","required":True},
        {"fieldId":"regions","fieldName":"Regions","type":"multiselect_blocks","options":["EU","MEA"]},
    ],"dashboardWidgets":[{"title":"Broken","fieldId":"retired"}],
    "scoringCriteria":[{"fieldId":"retired","weight":1}],
    "suppliers":[
        {"supplierId":"A","supplierName":"Alpha","customFields":{"rating":6,"regions":["EU","UNKNOWN"],"legacy":"old"}},
        {"supplierId":"B","supplierName":"Beta","customFields":{}},
    ]}
    findings=inspect_panel(core)
    codes=[x["code"] for x in findings]
    assert "broken_widget" in codes
    assert "broken_score" in codes
    assert "invalid_value" in codes
    assert "invalid_region" in codes
    assert "orphan_value" in codes
    assert "missing_required" in codes
    assert "missing_review" in codes
    assert inspect_panel({"supplierFields":[],"suppliers":[]})==[]


def test_dev046_data_quality_page_read_only(client):
    create_panel(client,"CMP-QUALITY")
    before=client.get("/api/panels/CMP-QUALITY").get_json()
    page=client.get("/data-quality")
    assert page.status_code==200
    assert b"Data quality dashboard" in page.data
    assert b"CMP-QUALITY" in page.data
    assert b"Data Quality" in client.get("/").data
    assert client.get("/api/panels/CMP-QUALITY").get_json()==before

def test_dev046_review_date_validation_is_configurable():
    from data_quality import inspect_panel
    core={"supplierFields":[],"suppliers":[
        {"supplierId":"A","supplierName":"Supplier A","customFields":{"expiry":"2026-13-01"}},
        {"supplierId":"B","supplierName":"Supplier B","customFields":{"expiry":"2026-11-01"}},
    ]}
    findings=inspect_panel(core,"expiry")
    assert any(f["code"]=="invalid_review" and f["supplierId"]=="A" for f in findings)
    assert not any(f["supplierId"]=="B" and f["code"] in ("missing_review","invalid_review") for f in findings)

def test_dev047_master_sync_preview_confirmation_and_conflict(client):
    import sqlite3
    from itsdangerous import URLSafeTimedSerializer
    create_panel(client, "CMP-MASTER-SYNC")
    assert client.post("/panels/CMP-MASTER-SYNC/suppliers/new",data={
        "supplier_id":"BP-123","supplier_name":"Panel name","address":"Panel address",
        "post_code":"ST1","custom_rating":"4"
    }).status_code==302
    with sqlite3.connect(client.application.config["DATABASE"]) as db:
        db.execute("INSERT OR REPLACE INTO supplier_master(bpid,supplier_name,address,post_code,address_source,active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                   ("BP-123","New master name","Master address","ST2","imported",1,"2026-01-01","2026-01-01"))
    preview=client.get("/panels/CMP-MASTER-SYNC/suppliers/master-sync")
    assert preview.status_code==200
    assert b"New master name" in preview.data
    assert b"Panel name" in preview.data
    before=client.get("/api/panels/CMP-MASTER-SYNC").get_json()["panel"]
    assert before["suppliers"][0]["supplierName"]=="Panel name"
    token=URLSafeTimedSerializer(client.application.secret_key,salt="catmanager-master-sync").dumps({
        "panel":"CMP-MASTER-SYNC",
        "changes":[{"supplierId":"BP-123","fields":{
            "supplierName":{"before":"Panel name","after":"New master name"},
            "address":{"before":"Panel address","after":"Master address"},
            "postCode":{"before":"ST1","after":"ST2"}}}]
    })
    with sqlite3.connect(client.application.config["DATABASE"]) as db:
        db.execute("UPDATE supplier_master SET supplier_name=? WHERE bpid=?",("Updated again","BP-123"))
    stale=client.post("/panels/CMP-MASTER-SYNC/suppliers/master-sync/confirm",data={"token":token},follow_redirects=True)
    assert b"not applied" in stale.data
    assert client.get("/api/panels/CMP-MASTER-SYNC").get_json()["panel"]["suppliers"][0]["supplierName"]=="Panel name"
    with sqlite3.connect(client.application.config["DATABASE"]) as db:
        db.execute("UPDATE supplier_master SET supplier_name=? WHERE bpid=?",("New master name","BP-123"))
    applied=client.post("/panels/CMP-MASTER-SYNC/suppliers/master-sync/confirm",data={"token":token})
    assert applied.status_code==302
    after=client.get("/api/panels/CMP-MASTER-SYNC").get_json()["panel"]
    assert after["suppliers"][0]["supplierName"]=="New master name"
    assert after["suppliers"][0]["address"]=="Master address"
    assert after["suppliers"][0]["customFields"]==before["suppliers"][0]["customFields"]
    assert any(e["action"]=="supplier_master_synced" for e in after["metadata"]["auditTrail"])
    assert client.post("/panels/CMP-MASTER-SYNC/suppliers/master-sync/confirm",data={"token":"bad"}).status_code==302
    assert client.get("/panels/UNKNOWN/suppliers/master-sync").status_code==404

def test_dev048_csrf_opt_in_blocks_untrusted_posts_and_embeds_form_tokens(tmp_path):
    import re
    secure=create_app({"TESTING":True,"SECRET_KEY":"secure-test-key","DATABASE":str(tmp_path/"csrf.sqlite"),
                       "SEED_DEMO":False,"CSRF_ENABLED":True})
    with secure.test_client() as c:
        page=c.get("/panels/new")
        assert page.status_code==200
        match=re.search(rb'name="_csrf_token" value="([^"]+)"',page.data)
        assert match
        body={"panel_id":"CSRF-PANEL","panel_name":"Secure Test","category":"Transformers",
              "business":"GI","region_level":"Country","region_value":"UK","owner":"A",
              "mdf_codes":["MDF-TR-001"],"lead_mdf_code":"MDF-TR-001","fields_json":"[]"}
        assert c.post("/panels/new",data=body).status_code==400
        assert c.post("/panels/new",data={**body,"_csrf_token":match.group(1).decode()}).status_code==302


def test_dev048_explicit_production_secret_guard_and_migration_ledger(tmp_path):
    import sqlite3
    with pytest.raises(RuntimeError,match="CATMANAGER_SECRET_KEY"):
        create_app({"REQUIRE_STRONG_SECRET":True,"SECRET_KEY":"dev-change-me",
                    "DATABASE":str(tmp_path/"reject.sqlite")})
    secure=create_app({"TESTING":True,"REQUIRE_STRONG_SECRET":True,"SECRET_KEY":"real-secret-key",
                       "DATABASE":str(tmp_path/"schema.sqlite"),"SEED_DEMO":False})
    with sqlite3.connect(secure.config["DATABASE"]) as db:
        rows=db.execute("SELECT version,description FROM schema_migrations").fetchall()
        assert rows==[(1,"Baseline CatManager schema")]
    secure2=create_app({"TESTING":True,"SECRET_KEY":"real-secret-key",
                        "DATABASE":secure.config["DATABASE"],"SEED_DEMO":False})
    with sqlite3.connect(secure2.config["DATABASE"]) as db:
        assert db.execute("SELECT COUNT(*) FROM schema_migrations").fetchone()[0]==1

def test_bug049_kpi_preview_works_with_csrf_enabled(tmp_path):
    import re
    secure=create_app({"TESTING":True, "DATABASE":str(tmp_path/"widget-csrf.sqlite"),
                       "SEED_DEMO":False,"SECRET_KEY":"preview-test-secret","CSRF_ENABLED":True})
    with secure.test_client() as c:
        form=c.get("/panels/new")
        token=re.search(rb'name="_csrf_token" value="([^"]+)"',form.data).group(1).decode()
        create_data={"panel_id":"CMP-PREVIEW","panel_name":"Preview panel","category":"Transformers",
                     "business":"GI","region_level":"Country","region_value":"UK",
                     "owner":"Test","mdf_codes":["MDF-TR-001"],"lead_mdf_code":"MDF-TR-001",
                     "fields_json":"[]","_csrf_token":token}
        assert c.post("/panels/new",data=create_data).status_code==302
        page=c.get("/panels/CMP-PREVIEW/configuration")
        assert page.status_code==200
        assert b'data-csrf-token="' in page.data
        preview_url="/panels/CMP-PREVIEW/widgets/preview"
        widgets={"widgets":[{"title":"Supplier Count","metric":"count","fieldId":"","format":"number","display":"panel"}]}
        blocked=c.post(preview_url,json=widgets)
        assert blocked.status_code==400
        allowed=c.post(preview_url,json=widgets,headers={"X-CSRF-Token":token})
        assert allowed.status_code==200
        assert allowed.is_json
        assert allowed.get_json()["widgets"][0]["title"]=="Supplier Count"


def test_bug049_widget_js_sends_csrf_and_checks_content_type():
    from pathlib import Path
    source=(Path(__file__).resolve().parents[1]/"static"/"app.js").read_text()
    assert '"X-CSRF-Token":widgetRoot.dataset.csrfToken' in source
    assert 'contentType.includes("application/json")' in source

def test_ux051_supplier_tab_default_and_secondary_views(client):
    create_panel(client,"CMP-UX051")
    default=client.get("/panels/CMP-UX051")
    assert default.status_code==200
    html=default.get_data(as_text=True)
    assert 'aria-current="page">Suppliers' in html
    assert '<h2>Suppliers</h2>' in html
    assert '<h2>Supplier coverage and sourcing concentration</h2>' not in html
    assert 'More actions' in html
    for name in ("Edit Panel","Panel Configuration","Add Supplier",
                 "Management Excel Report","Sync Supplier Master","Save Template"):
        assert name in html
    for tab, heading in (
        ("overview","Supplier coverage and sourcing concentration"),
        ("actions","Qualification review alerts"),
        ("history","Panel activity history"),
    ):
        response=client.get("/panels/CMP-UX051?tab="+tab)
        assert response.status_code==200
        assert heading in response.get_data(as_text=True)
    assert b'<h2>Suppliers</h2>' in client.get("/panels/CMP-UX051?tab=unrecognised").data
