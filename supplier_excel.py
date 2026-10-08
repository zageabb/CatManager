"""Supplier XLSX exchange with strict, value-only parsing."""
import hashlib
import io
import json
from datetime import date, datetime

from openpyxl import Workbook, load_workbook


def signature(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def columns(fields):
    return ["Supplier ID", "Baseline"] + [f["fieldId"] for f in fields]


def encode_value(value, kind):
    if kind == "multiselect_blocks":
        return json.dumps(value if isinstance(value, list) else [])
    if kind == "boolean":
        return "Yes" if value is True else "No" if value is False else ""
    return "" if value is None else value


def export_suppliers(core):
    fields = core.get("supplierFields", [])
    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    ws.append(columns(fields))
    for supplier in core.get("suppliers", []):
        custom = supplier.get("customFields", {})
        ws.append([supplier["supplierId"], signature(custom)] +
                  [encode_value(custom.get(f["fieldId"]), f["type"]) for f in fields])
    # Force text fields to remain literal strings, even when they start with
    # '=', '+', '-' or '@'. This prevents exported supplier data becoming formulas.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].hidden = True
    for i, field in enumerate(fields, start=3):
        ws.column_dimensions[ws.cell(1,i).column_letter].width = 23
        ws.cell(1,i).comment = None
    note = wb.create_sheet("Instructions")
    for line in [
        "CatManager Supplier Excel Exchange",
        "Edit only custom-field values in Suppliers. Do not modify Supplier ID, Baseline or headers.",
        "Multi-select options: enter a JSON list, e.g. [\"EU\", \"MEA\"].",
        "Boolean values: Yes, No or blank. Stars: 0–5 in 0.5 increments.",
        "Import shows a preview and requires confirmation. Changed source records are rejected.",
        "New and deleted suppliers are not supported in this workflow.",
    ]:
        note.append([line])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def parse_value(raw, field):
    kind = field["type"]
    if isinstance(raw, (date, datetime)):
        raw = raw.date().isoformat() if isinstance(raw,datetime) else raw.isoformat()
    val = "" if raw is None else str(raw).strip()
    if kind == "multiselect_blocks":
        try:
            items = json.loads(val or "[]")
        except ValueError as exc:
            raise ValueError("Expected a JSON list") from exc
        if not isinstance(items, list) or any(not isinstance(x,str) or x not in field.get("options",[]) for x in items):
            raise ValueError("Invalid selected options")
        return list(dict.fromkeys(items))
    if kind == "boolean":
        if not val:
            return None
        if val.casefold() in ("yes","true","1"):
            return True
        if val.casefold() in ("no","false","0"):
            return False
        raise ValueError("Expected Yes, No or blank")
    if kind == "dropdown" and val and val not in field.get("options",[]):
        raise ValueError("Option not configured")
    if kind in ("number","stars") and val:
        try:
            number=float(val)
        except ValueError as exc:
            raise ValueError("Expected a number") from exc
        if not (-1e100 < number < 1e100):
            raise ValueError("Invalid number")
        if kind=="stars" and not (0<=number<=5 and number*2==int(number*2)):
            raise ValueError("Stars must be 0–5 in half-point steps")
        return number
    if field.get("required") and not val:
        raise ValueError("Required field")
    return val


def preview_import(content, core):
    if len(content)>5_000_000:
        raise ValueError("File exceeds 5 MB")
    try:
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        ws=wb["Suppliers"]
    except Exception as exc:
        raise ValueError("Invalid Excel workbook or missing Suppliers tab") from exc
    fields=core.get("supplierFields",[])
    expected=columns(fields)
    rows=ws.iter_rows(values_only=True)
    if list(next(rows, []))!=expected:
        raise ValueError("Column headers or custom-field configuration do not match this panel")
    existing={s["supplierId"]:s for s in core.get("suppliers",[])}
    changes=[]
    seen=set()
    for row in rows:
        if not any(v is not None for v in row):
            continue
        if len(seen)>=5000:
            raise ValueError("Too many supplier rows")
        if len(row)!=len(expected):
            raise ValueError("Incorrect number of columns")
        supplier_id=str(row[0] or "").strip()
        if supplier_id in seen or supplier_id not in existing:
            raise ValueError("Unknown or duplicate supplier ID: "+supplier_id)
        seen.add(supplier_id)
        supplier=existing[supplier_id]
        before=supplier.get("customFields",{})
        if row[1]!=signature(before):
            raise ValueError("Supplier "+supplier_id+" changed since workbook export")
        new=dict(before)
        for field, raw in zip(fields,row[2:]):
            try:
                new[field["fieldId"]]=parse_value(raw,field)
            except ValueError as exc:
                raise ValueError(supplier_id+" / "+field["fieldName"]+": "+str(exc)) from exc
        if new!=before:
            changes.append({"supplierId":supplier_id,"before":before,"after":new})
    if seen!=set(existing):
        raise ValueError("Workbook is missing one or more suppliers")
    return changes
