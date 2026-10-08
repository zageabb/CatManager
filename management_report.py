"""Generate a read-only Excel management report from current panel calculations."""
from io import BytesIO
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def build_management_workbook(core, scores, coverage, alerts, qualification_field):
    wb=Workbook()
    overview=wb.active
    overview.title="Overview"
    overview.append(["CatManager — Supplier Management Summary"])
    overview.append(["Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    overview.append(["Panel ID",core.get("panelId","")])
    overview.append(["Panel",core.get("panelName","")])
    overview.append(["Category",core.get("category","")])
    overview.append(["Suppliers",len(core.get("suppliers",[]))])
    overview.append(["Lead MDF",coverage["leadMdfCode"]])
    overview.append(["MDF scope",", ".join(coverage["mdfCodes"])])
    overview.append(["Review overdue",sum(a["reviewStatus"]=="Overdue" for a in alerts)])
    overview.append(["Reviews due soon",sum(a["reviewStatus"]=="Due soon" for a in alerts)])
    overview.append(["Uncovered regions",sum(r["risk"]=="Uncovered" for g in coverage["groups"] for r in g["regions"])])
    overview.append(["Single-source regions",sum(r["risk"]=="Single source" for g in coverage["groups"] for r in g["regions"])])
    overview.append(["Note","Coverage represents selected regions, not verified supplier-to-MDF capability."])
    score_index={x["supplierId"]:x for x in scores}
    alert_index={x["supplierId"]:x for x in alerts}
    ws=wb.create_sheet("Suppliers")
    ws.append(["Supplier ID","Supplier Name","Qualification","Score / 5","Scoring status","Review date","Review status","Open actions"])
    for supplier in core.get("suppliers",[]):
        sid=supplier.get("supplierId","")
        score=score_index.get(sid,{})
        alert=alert_index.get(sid,{})
        ws.append([sid,supplier.get("supplierName",""),
                   str(supplier.get("customFields",{}).get(qualification_field,"") or ""),
                   score.get("score"),"Complete" if score.get("complete") else "Incomplete / not configured",
                   str(alert.get("reviewDate") or ""),alert.get("reviewStatus",""),
                   sum(x.get("status")!="Completed" for x in supplier.get("actions",[]))])
    regions=wb.create_sheet("Regional coverage")
    regions.append(["Field","Region","Supplier count","Qualified count","Risk","Supplier IDs"])
    for group in coverage["groups"]:
        for row in group["regions"]:
            regions.append([group["fieldName"],row["region"],row["count"],row["qualifiedCount"],row["risk"],", ".join(row["supplierIds"])])
    actions=wb.create_sheet("Supplier actions")
    actions.append(["Supplier ID","Supplier name","Action","Owner","Due date","Status"])
    for supplier in core.get("suppliers",[]):
        for item in supplier.get("actions",[]):
            if item.get("status")=="Completed": continue
            actions.append([supplier.get("supplierId",""),supplier.get("supplierName",""),
                            item.get("title",""),item.get("owner",""),item.get("dueDate",""),item.get("status","")])
    for page in wb:
        page.freeze_panes="A2"
        if page is not overview:
            page.auto_filter.ref=page.dimensions
        for cell in page[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="29364B")
        for row in page:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type="s"  # avoid treating supplier-entered text as formulas
                cell.alignment=Alignment(vertical="top",wrap_text=True)
        for col in page.columns:
            letter=get_column_letter(col[0].column)
            page.column_dimensions[letter].width=min(52,max(16,max(len(str(cell.value or "")) for cell in col)+2))
    buf=BytesIO()
    wb.save(buf)
    return buf.getvalue()
